"""Read-only verification of hashes, template rows and final-only evidence."""
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parent
TYPES=('A_uav','B_uav','C_uav','A_battery','B_battery','C_battery','R_uav','R_component')
def read(p):
    with p.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def main():
    manifest=json.loads((ROOT/'manifest.json').read_text(encoding='utf-8'))
    for f in manifest['files']:
        p=ROOT/f['path'];assert p.stat().st_size==f['bytes'] and hashlib.sha256(p.read_bytes()).hexdigest()==f['sha256'],f['path']
    out=ROOT/'results/q4_final';src=ROOT/'results/q3_overlay'
    d=json.loads((out/'q4_results.json').read_text(encoding='utf-8'))
    audit=json.loads((ROOT/'results/q4_optimality_audit/optimality_certificate.json').read_text(encoding='utf-8'))
    assert audit['status']=='PASS' and audit['nonempty_subsets_checked']==255
    w=load_workbook(out/'Q4_结果提交.xlsx');original=load_workbook(src/'Q3_结果提交.xlsx')
    for s in original:
        if s.title=='Q4_分区配置':continue
        a,b=list(s.values),list(w[s.title].values);assert len(a)==len(b)
        for ra,rb in zip(a,b):
            for va,vb in zip(ra,rb):
                if isinstance(va,(int,float)) and isinstance(vb,(int,float)):assert abs(va-vb)<=1e-9*max(1,abs(va))
                else:assert va==vb
    sh=w['Q4_分区配置'];assert list(sh.values)[0]==list(original[sh.title].values)[0]
    actual=[tuple(sh.cell(i,j).value for j in range(1,12)) for i in range(2,sh.max_row+1) if sh.cell(i,1).value is not None]
    expected=[]
    for k,units,gap in (('2',29,2),('3',34,4)):
        x=d['final_solutions'][k];z=audit['results'][k]
        assert x==d['pooled_results'][k]['recommended']
        assert x['total_units']==z['global_minima']['units']==units
        assert x['gap_units']==z['global_minima']['gap']==gap
        assert z['partitions']=={'2':127,'3':966}[k]
        for gi,g in enumerate(x['groups'],1):expected.append((int(k),f'G{gi}',';'.join(g['sites']),*[g['resources'][t] for t in TYPES]))
    assert actual==expected and len(actual)==5
    assignments=read(out/'组内实体映射.csv');comm=read(out/'通信映射.csv');boxes=read(out/'货箱归属.csv')
    source_comm=read(src/'continuous_communication_certificate.csv')
    for k in ('2','3'):
        c=[r for r in comm if r['K']==k];b=[r for r in boxes if r['K']==k]
        assert len(c)==296 and len(b)==len({r['box_id'] for r in b})==80
        assert len({(r['运输架次'],r['t0_s'],r['t1_s']) for r in c})==296
        gids={p:str(i) for i,g in enumerate(d['final_solutions'][k]['groups'],1) for p in g['sorties']}
        for r,s in zip(sorted(c,key=lambda x:(int(x['运输架次']),float(x['t0_s']))),sorted(source_comm,key=lambda x:(int(x['route']),float(x['t0'])))):
            assert r['任务组']=='G'+gids[int(s['route'])]
            assert int(r['运输架次'])==int(s['route']) and abs(float(r['t0_s'])-float(s['t0']))<1e-7 and abs(float(r['t1_s'])-float(s['t1']))<1e-7
            assert abs(float(r['margin_dB'])-float(s['margin_dB']))<1e-7
            if s['provider']=='G01':assert r['本组通信提供者']=='G01'
            else:
                match=[v for v in assignments if v['K']==k and v['group']==r['任务组'][1:] and v['resource_type']=='R_uav' and v['q3_task']==r['原中继任务']]
                assert len(match)==1 and r['本组通信提供者']==match[0]['local_resource']
        assert all(r['variant']=='recommended' for r in assignments if r['K']==k)
    errors={'#VALUE!','#DIV/0!','#REF!','#NAME?','#NULL!','#NUM!','#N/A'}
    assert not any(c.data_type=='e' or (isinstance(c.value,str) and c.value in errors) for s in w for row in s for c in row)
    print('PASS: hashes, 5 template rows, preserved Q1-Q3, 160 box mappings, 592 communication intervals, global minima 29/34 and shortage 2/4')
if __name__=='__main__':main()
