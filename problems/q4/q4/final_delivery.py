"""Publish the Q4 minimum-shortage/minimum-size solutions in the official template."""
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
import solve_overlay as model
import verify_overlay as verifier

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'results/q4_final'
NAMES=('A型运输机','B型运输机','C型运输机','A型电池','B型电池','C型电池','中继机','能源组件')

def csv_write(name, rows):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        csv.writer(f).writerows(rows)

def compact(x):
    r=x['total']
    return (sum(r[t+'_uav'] for t in 'ABC'),sum(r[t+'_battery'] for t in 'ABC'),r['R_uav'],r['R_component'])

def same_sheet(a,b):
    assert a.max_row==b.max_row and a.max_column==b.max_column
    for ra,rb in zip(a,b):
        for ca,cb in zip(ra,rb):
            if isinstance(ca.value,(int,float)) and isinstance(cb.value,(int,float)):
                assert abs(ca.value-cb.value)<=1e-9*max(1,abs(ca.value))
            else: assert ca.value==cb.value

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    ts,rs,cert,boxes=model.inputs()
    blocks=model.base.connected_components(ts); rm=model.base.relay_route_map(rs,cert)
    d=dict(components=blocks,inventory=model.STOCK,results={},pooled_results={},target_comparison={},final_solutions={})
    header=['gap_units','total_units','work_cv','relay_task_copies','groups',*['need_'+t for t in model.TYPES]]
    for k in (2,3):
        for pooled,key in ((False,'results'),(True,'pooled_results')):
            info,rows=model.search(ts,rs,rm,boxes,blocks,k,pooled)
            if pooled:
                info['recommended']=info['resource_first'];info['balance_cap']=None
                info['selection']='minimum inventory shortage, then configuration units, then CV'
                x=info['recommended']; d['final_solutions'][str(k)]=x
                d['target_comparison'][str(k)]={'selected':x,'resource_first':x}
                assert x['total_units']==min(r['total_units'] for r in rows)
                assert x['gap_units']==min(r['gap_units'] for r in rows)
            d[key][str(k)]=info
            csv_write(f'q4_{k}group'+('_pooled' if pooled else '')+'_search.csv',[header]+[
                [x['gap_units'],x['total_units'],x['work_cv'],x['relay_task_copies'],
                 ' | '.join(';'.join(g['sites']) for g in x['groups']),*[x['total'][t] for t in model.TYPES]] for x in rows])
    names=['transport_sorties.csv','relay_sorties.csv','continuous_communication_certificate.csv',
           'box_deliveries.csv','q3_final_solution.mat','q3_metrics.csv','validator_PASS.txt','问题三_优化结果.md']
    d['q3_source']='results/q3_overlay'
    d['source_sha256']={n:hashlib.sha256((model.SOURCE/n).read_bytes()).hexdigest() for n in names}
    d['resource_policy']='Frozen task geometry, times and communication relations; same-type entities may be reassigned within each group'
    d['schema_note']='final_solutions and pooled_results.recommended are authoritative; results contains original-ID comparison evidence only'
    (OUT/'q4_results.json').write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
    verifier.OUT=OUT; verifier.main()
    assert [compact(d['final_solutions'][k]) for k in ('2','3')]==[(9,15,2,3),(9,15,4,6)]
    publish(d,ts,rs,cert,boxes)
    print('Q4 final: 2 groups 9/15/2/3 (29 units, shortage 2); 3 groups 9/15/4/6 (34 units, shortage 4)')

def publish(d,ts,rs,cert,boxes):
    group_rows=[['K','任务组编号','服务区列表',*NAMES,'箱数','运输架次','中继架次','运输工作量h','中继工作量h','总工作量h','最晚返航min','运输能耗kWh','中继能耗kWh']]
    inventory=[['K','资源类型','现有库存','独立组需求','配置缺口','峰值下界','超出峰值的配置冗余']]
    summary=[['K','运输机','电池','中继机','能源组件','总件数','总缺口','工作量CV','最大最小工作量比','运输架次','中继架次','新增中继架次','总能耗kWh','最晚返航min']]
    communication=[['K','任务组','运输架次','t0_s','t1_s','原中继任务','本组通信提供者','margin_dB']]
    deliveries=[['K','任务组','box_id','area','sortie_id']]
    for k in ('2','3'):
        x=d['final_solutions'][k]
        for gi,g in enumerate(x['groups'],1):
            group_rows.append([int(k),f'G{gi}',';'.join(g['sites']),*[g['resources'][t] for t in model.TYPES],g['boxes'],len(g['sorties']),len(g['relay_ids']),g['transport_work_s']/3600,g['relay_work_s']/3600,g['work_s']/3600,g['makespan_s']/60,g['transport_energy_kWh'],g['relay_energy_kWh']])
            for b in boxes:
                if b['area'] in g['sites']: deliveries.append([int(k),f'G{gi}',b['box_id'],b['area'],b['sortie_id']])
            for c in cert:
                if int(c['route']) not in g['sorties']: continue
                task=''; provider='G01'
                if c['provider']!='G01':
                    hits=[r for r in rs if r['uav']==c['provider'] and float(r['serviceStart_s'])<=float(c['t0'])+1e-6 and float(r['serviceEnd_s'])>=float(c['t1'])-1e-6]
                    assert len(hits)==1;task=hits[0]['id']
                    provider=f"K{k}-recommended-G{gi}-R_uav-{g['pooled_assignment']['R_uav'][task]}"
                communication.append([int(k),f'G{gi}',int(c['route']),float(c['t0']),float(c['t1']),task,provider,float(c['margin_dB'])])
        for t,label in zip(model.TYPES,NAMES):
            inventory.append([int(k),label,model.STOCK[t],x['total'][t],x['gap'][t],x['peak_lower_bound'][t],x['redundancy'][t]])
        summary.append([int(k),*compact(x),x['total_units'],x['gap_units'],x['work_cv'],x['work_ratio'],23,x['relay_task_copies'],x['relay_task_copies']-4,sum(g['transport_energy_kWh']+g['relay_energy_kWh'] for g in x['groups']),max(g['makespan_s'] for g in x['groups'])/60])
    with (OUT/'q4_pooled_resource_assignments.csv').open(encoding='utf-8-sig',newline='') as f:
        all_resource=list(csv.reader(f))
    resource=[all_resource[0]]+[r for r in all_resource[1:] if r[1]=='recommended']
    sheets=[('Q4_方案汇总','方案汇总.csv',summary),('Q4_逐组明细','逐组配置.csv',group_rows),
            ('Q4_库存缺口冗余','库存缺口冗余.csv',inventory),('Q4_组内实体映射','组内实体映射.csv',resource),
            ('Q4_通信映射','通信映射.csv',communication),('Q4_货箱归属','货箱归属.csv',deliveries)]
    for _,name,rows in sheets: csv_write(name,rows)
    w=load_workbook(model.SOURCE/'Q3_结果提交.xlsx')
    template=w['Q4_分区配置']
    for row in template.iter_rows(min_row=2):
        for cell in row: cell.value=None
    for i,row in enumerate(group_rows[1:],2):
        for j,value in enumerate(row[:11],1): template.cell(i,j,value)
    for title,_,rows in sheets:
        s=w.create_sheet(title)
        for row in rows:s.append(row)
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='30576D');c.alignment=Alignment(wrap_text=True)
        s.row_dimensions[1].height=32
        for col in s.columns:
            width=min(55,max(14,max(len(str(c.value or '')) for c in col)+2))
            s.column_dimensions[col[0].column_letter].width=width
    path=OUT/'Q4_结果提交.xlsx';w.save(path)
    check=load_workbook(path);source=load_workbook(model.SOURCE/'Q3_结果提交.xlsx')
    for s in source:
        if s.title!='Q4_分区配置':same_sheet(s,check[s.title])
    assert list(check['Q4_分区配置'].values)[0]==list(source['Q4_分区配置'].values)[0]
    for i,row in enumerate(group_rows[1:],2): assert [check['Q4_分区配置'].cell(i,j).value for j in range(1,12)]==row[:11]
    assert sum(r[0]==2 for r in communication[1:])==296 and sum(r[0]==3 for r in communication[1:])==296
    assert sum(r[0]==2 for r in deliveries[1:])==80 and sum(r[0]==3 for r in deliveries[1:])==80
    report(d,group_rows,inventory,summary)

def report(d,groups,inventory,summary):
    lines=['# 问题四：任务分区与资源配置最终结果','',
           '本次主方案采用总件数与库存缺口同时达到全局最少的结果：两组9/15/2/3，三组9/15/4/6。数字依次表示运输机、电池、中继机、能源组件。所有最优性结论限定于本次固定的问题三任务、时刻及通信保障关系，允许组内同型资源重新指派，不允许跨组共享。','',
           '## 题目主要需求与处理方式','',
           '|题目需求|本方案处理|交付证据|','|---|---|---|',
           '|15个服务区划为2组、3组，每区唯一归属且组非空|共访关系形成8个不可拆分任务块，穷举127/966种分区|模板Q4主表、全部搜索CSV|',
           '|保持Q3分箱、访问顺序、运输与中继任务安排、通信保障关系|冻结轨迹和全部任务时刻；同一运输架次所访区同组；跨组使用的中继任务整段复制|Q3源输入哈希、货箱归属、通信映射|',
           '|各组独立核算四类资源，不跨组调配|时间DAG最大匹配求最少路径覆盖，构造组内实体编号|组内实体映射、独立核验|',
           '|比较配置规模、冗余、工作量均衡、库存缺口|逐型号统计，公开工作量CV和中继复制代价|方案汇总、库存缺口冗余、逐组配置|',
           '|超库存时给出缺口和原因|按型号正差计缺口，不用异型库存抵扣|两组缺2件、三组缺4件|','',
           '## 模型与求解','',
           '先最小化各型号库存缺口之和，再最小化配置总件数，再最小化工作量CV。此次选解恰好同时达到前两项各自的全局下界，而非仅在缺口最小时求得条件最少件数。件数不是采购成本；题目未规定该优先顺序或CV阈值，此处公开采用资源优先偏好。','',
           '共访闭包为：'+'；'.join('、'.join(b) for b in d['components'])+'。全部无标号非空分区已穷举。固定每组任务后，对8种资源分别构造时间有向无环图，以任务数减最大二分匹配数获得最少实体数；每条路径分配一个组内实体。','',
           '运输机衔接检查前次返航≤后次装载开始，以及前次起飞≤后次固定预准备开始。电池再次使用前须按两阶段充电模型充至100%；中继机计入周转，组件计入满电等待。该预准备解释继承Q3；若冻结原实体ID或改变预准备规则，应重新求解，不能直接沿用本结果。','',
           '资源不足不作为分区不可行条件，因为原题要求评估并报告缺口。以下配置必须补足所列设备后才能独立执行。固定网关G01为共同基础设施，不属于四类机动资源。','',
           '## 分区与配置','',
           '|组数|组|服务区|运输机A/B/C|电池A/B/C|中继机|组件|箱数|运输/中继架次|工作量h|',
           '|---|---|---|---|---|---:|---:|---:|---|---:|']
    for row in groups[1:]:
        k,gi,sites,*_=row
        lines.append(f"|{k}|{gi}|{sites.replace(';','、')}|{'/'.join(map(str,row[3:6]))}|{'/'.join(map(str,row[6:9]))}|{row[9]}|{row[10]}|{row[11]}|{row[12]}/{row[13]}|{row[16]:.6f}|")
    lines+=['','## 规模、均衡与中继复制代价','','|组数|四类配置|总件数|缺口|CV|工作量最大/最小|中继架次|总能耗kWh|最晚返航min|','|---|---|---:|---:|---:|---:|---:|---:|---:|']
    for row in summary[1:]:
        lines.append(f"|{row[0]}|{'/'.join(map(str,row[1:5]))}|{row[5]}|{row[6]}|{row[7]:.9f}|{row[8]:.6f}|{row[10]}|{row[12]:.9f}|{row[13]:.9f}|")
    lines+=['','工作量定义为本组运输operation_s总和与中继任务start至return时长之和；它是任务工作量指标，不等于总准备工时或联合完工时间。CV为组工作量总体标准差除以均值。两组方案将S006单独成组，均衡较差；三组CV略有改善，但仍不是均衡最优。','',
            '问题三原有23个运输架次、4个中继架次。两组无需复制额外中继架次；三组将跨组保障的中继任务复制，共7个中继架次，增加3架次及相应能耗。运输架次及运输能耗保持不变；时刻冻结，两方案联合最晚返航相同。','',
            '## 库存缺口与资源冗余','','|资源|库存|两组需求|两组缺口|三组需求|三组缺口|','|---|---:|---:|---:|---:|---:|']
    a,b=[d['final_solutions'][k] for k in ('2','3')]
    for t,label in zip(model.TYPES,NAMES):lines.append(f"|{label}|{model.STOCK[t]}|{a['total'][t]}|{a['gap'][t]}|{b['total'][t]}|{b['gap'][t]}|")
    lines+=['','两组须补充1架C型运输机和1组C型电池；三组在此基础上再补充2架中继机。原因是原Q3跨任务复用的实体不能跨组调配，以及原中继任务跨组保障时必须各组配置独立副本。附件组件库存为6组，三组方案恰好用满；不能把截图中的3组当作库存。','',
            '冗余分两个口径：①组内配置数减资源占用并发峰值下界，两个主方案8种资源均为0；且最大匹配已证明各组资源数最少，不存在可删除的闲置配置。②分组独立执行造成的中继任务复制，三组比两组多3个中继架次、多2架中继机和3组组件。前一种为0不代表没有独立分组带来的额外代价。','',
            '## 最优性与取舍','',
            '独立程序不导入求解器，重新建立共访块、重算255个非空任务块子集，并用SciPy匹配与独立标签枚举复核全部1093个分区。两组最低缺口2（2个分区），最低总件数29（唯一分区）；三组最低缺口4（唯一分区），最低总件数34（2个分区）。忽略任务组名称排列。','',
            '两组和三组的均衡单目标最低CV分别约0.002136、0.012267，但对应配置分别为13/17/4/6和14/17/6/9，缺口10和16。本次资源优先方案因此不是所有指标同时最优。三组主方案比两组多5件配置和2件缺口，换取第三个独立执行单元，不能宣称显著改善均衡。','',
            '本报告不证明Q3本身全局最优。若重设计Q3任务结构、时刻、通信关联或坚持原实体ID，最优性证书不再适用。沿用Q3连续通信物理证书，组内重指派与完整任务复制不改变链路端点和时段；不重新求解地形传播。','',
            '## 文件与复现','',
            '`Q4_结果提交.xlsx`在原Q3工作簿副本中填入官方`Q4_分区配置`主表，保留其他原有工作表；新增汇总、逐组明细、库存缺口冗余、实体、通信和货箱映射。每个方案包含15区、23个运输任务、80箱和296个连续通信区间。','',
            '仓库包根目录执行`python run_problem4.py`，再执行`python verify_bundle.py`。独立最优性证书位于`results/q4_optimality_audit/optimality_certificate.json`。详情及后续开发约束见包根README与`docs/GitHub任务指令.md`。','']
    (OUT/'问题四_最终结果与方法.md').write_text('\n'.join(lines),encoding='utf-8')

if __name__=='__main__': main()
