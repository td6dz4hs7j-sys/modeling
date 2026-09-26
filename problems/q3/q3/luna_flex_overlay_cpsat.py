import csv, json, math, time
from openpyxl import load_workbook
from pathlib import Path
from ortools.sat.python import cp_model
ROOT=Path(__file__).resolve().parents[1]
O=ROOT/'results'/'q3_overlay'
F=ROOT/'results'/'q3_fresh'

def readcsv(p):
    with open(p,newline='',encoding='utf-8-sig') as f:return list(csv.DictReader(f))
R=readcsv(F/'source_route_slack.csv')
P=readcsv(O/'selected_site_physics.csv')
W=readcsv(O/'continuous_demand_groups.csv')
assert len(R)==23 and len(P)==4
# Integer conservative data for CP-SAT.
route=[]
for k,r in enumerate(R):
    take=float(r['takeoff_s']);ret=float(r['return_s']);ready=float(r['battery_ready_s'])
    typ=str(r['battery_id']).split('-')[0]
    route.append({'route':k+1,'type':typ,'flight':math.ceil(ret-take-1e-9),
       'charge':math.ceil(ready-ret-1e-9),'load':math.ceil(float(r['load_duration_s'])-1e-9),
       'prep':math.ceil(float(r['prep_duration_s'])-1e-9),
       'upper':math.floor(take+float(r['expected_slack_s'])+1e-9),
       'expected_slack':float(r['expected_slack_s']),'orig':take,
       'uav_source':r['uav_id'],'battery_source':r['battery_id']})
assert {x['type'] for x in route}=={'A','B','C'}
model=cp_model.CpModel();H=max(max(x['upper']+x['flight'] for x in route),12000)+12000
# Transport takeoff variables: no source-Q2 lower bound; expected-deadline upper bound only.
t=[];transport_intervals=[];battery_intervals=[]
for i,r in enumerate(route):
    lo=max(r['prep'],r['load'])
    ti=model.new_int_var(lo,r['upper'],f'takeoff_{i+1}');t.append(ti)
    transport_intervals.append(model.new_interval_var(ti-r['load'],r['load']+r['flight'],ti+r['flight'],f'transport_{i+1}'))
    battery_intervals.append(model.new_interval_var(ti,r['flight']+r['charge'],ti+r['flight']+r['charge'],f'battery_{i+1}'))
for typ,cap in [('A',4),('B',2),('C',2)]:
    model.add_cumulative([transport_intervals[i] for i,r in enumerate(route) if r['type']==typ],
                         [1]*sum(r['type']==typ for r in route),cap)
for typ,cap in [('A',6),('B',4),('C',4)]:
    model.add_cumulative([battery_intervals[i] for i,r in enumerate(route) if r['type']==typ],
                         [1]*sum(r['type']==typ for r in route),cap)
# Relay mission intervals and two physical relay UAVs.
relay_start=[];relay_end=[];relay_uav=[];relay_occ=[]
PREP=180;LINK=30;TURN=300
for j,p in enumerate(P):
    out=math.ceil(float(p['outbound_s'])-1e-9);ret=math.ceil(float(p['return_s'])-1e-9)
    maxservice=math.floor(float(p['max_service_s'])+1e-9)
    earliest=PREP+out+LINK
    s=model.new_int_var(earliest,H,f'relay_service_start_{j+1}')
    e=model.new_int_var(earliest+1,H,f'relay_service_end_{j+1}')
    model.add(e>=s+1);model.add(e-s<=maxservice)
    relay_start.append(s);relay_end.append(e)
    u=model.new_bool_var(f'relay_R01_{j+1}');relay_uav.append(u)
    dur=PREP+out+LINK+ret+TURN
    # Occupancy [s-(prep+out+link), e+return+turn], with conservative ceil flight times.
    size_var=model.new_int_var(dur+1,H,f'relay_occ_size_{j+1}')
    model.add(size_var==e-s+dur)
    end_var=model.new_int_var(0,H,f'relay_occ_end_{j+1}')
    model.add(end_var==e+ret+TURN)
    relay_occ.append((s-earliest,size_var,end_var,u,j))
# Each continuous demand group is assigned to exactly one eligible relay point.
assign_by_site=[[] for _ in P]
for k,w in enumerate(W):
    ri=int(w['route'])-1;mask=int(w['site_mask']);choices=[]
    for j in range(4):
        if mask & (1<<j):
            z=model.new_bool_var(f'cover_{k}_{j}');choices.append(z);assign_by_site[j].append(z)
            start_bound=math.floor(float(w['start_offset_s'])-.01)
            end_bound=math.ceil(float(w['end_offset_s'])+.01)
            model.add(relay_start[j]<=t[ri]+start_bound).only_enforce_if(z)
            model.add(relay_end[j]>=t[ri]+end_bound).only_enforce_if(z)
    if not choices:raise ValueError(f'group {k} has empty site mask')
    model.add(sum(choices)==1)
for j in range(4):model.add(sum(assign_by_site[j])>=1)
# Two relay aircraft: each task has exactly one of R01/R02; optional occupancy intervals capacity one each.
# relays are fixed to 4 distinct initially charged components R-B01..R-B04.
relay_intervals_by_uav=[[],[]]
for s,size,e,u,j in relay_occ:
    model.add(s>=0)
    int0=model.new_optional_interval_var(s,size,e,u,f'R01_interval_{j+1}')
    not_u=u.Not()
    int1=model.new_optional_interval_var(s,size,e,not_u,f'R02_interval_{j+1}')
    relay_intervals_by_uav[0].append(int0);relay_intervals_by_uav[1].append(int1)
for arr in relay_intervals_by_uav:model.add_no_overlap(arr)
# Makespan includes final return of transport and relay tasks.
C=model.new_int_var(0,H,'Cmax')
for i,r in enumerate(route):model.add(C>=t[i]+r['flight'])
for j,p in enumerate(P):model.add(C>=relay_end[j]+math.ceil(float(p['return_s'])-1e-9))
model.minimize(C)
solver=cp_model.CpSolver();solver.parameters.max_time_in_seconds=90;solver.parameters.num_search_workers=8;solver.parameters.random_seed=20260926
start_time=time.time();status=solver.solve(model)
status_name=solver.status_name(status)
meta={'status':status_name,'wall_s':solver.wall_time,'branches':solver.num_branches,'objective_s':solver.value(C) if status in (cp_model.OPTIMAL,cp_model.FEASIBLE) else None,
      'limits':'23 Q2 routes/groups/order/type fixed; same-type UAV and batteries reassigned; takeoff can advance/delay within 0..per-route expected-deadline upper; UAV capacities A4/B2/C2; battery A6/B4/C4; four relay sites/tasks; two relay UAVs; distinct initial components; 130 demand groups; conservative integer seconds; no continuous certification'}
(O/'luna_flex_solver.json').write_text(json.dumps(meta,indent=2),encoding='utf-8')
print(meta,flush=True)
if status not in (cp_model.OPTIMAL,cp_model.FEASIBLE):raise SystemExit(0)
# Greedy interval coloring against the stated same-type inventory.
uav_caps={'A':4,'B':2,'C':2};battery_caps={'A':6,'B':4,'C':4}
route_times=[solver.value(x) for x in t]
def assign_pool(kind,caps,intervals):
    result={};
    if kind=='uav':
        wb=load_workbook(ROOT/'input'/'数据'/'无人机应急物资运输基础数据'/'运输无人机数据.xlsx',data_only=True,read_only=True);sh=wb['数据']
        inv={typ:[str(sh.cell(row,1).value) for row in range(9,17) if str(sh.cell(row,2).value)==typ] for typ in caps};wb.close()
    else: inv={typ:[f'{typ}-B{i:02d}' for i in range(1,cap+1)] for typ,cap in caps.items()}
    for typ,cap in caps.items():
        ids=inv[typ];free={x:0 for x in ids}
        order=sorted([i for i,r in enumerate(route) if r['type']==typ],key=lambda i:(route_times[i],i))
        for i in order:
            st,en=intervals(i);available=[x for x in ids if free[x]<=st]
            if not available:raise RuntimeError(f'greedy {kind} coloring failed type={typ} route={i+1} start={st}')
            ident=min(available,key=lambda x:free[x]);result[i]=ident;free[ident]=en
    return result
uav_ids=assign_pool('uav',uav_caps,lambda i:(route_times[i]-route[i]['load'],route_times[i]+route[i]['flight']))
battery_ids=assign_pool('battery',battery_caps,lambda i:(route_times[i],route_times[i]+route[i]['flight']+route[i]['charge']))
with open(O/'luna_flex_transport.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f);w.writerow(['route','takeoff_s','uav_id','battery_id'])
    for i,r in enumerate(route):w.writerow([i+1,route_times[i],uav_ids[i],battery_ids[i]])
with open(O/'luna_flex_relays.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f);w.writerow(['site_row','uav','serviceStart_s','serviceEnd_s','component_id'])
    for j in range(4):w.writerow([j+1,'R01' if solver.value(relay_uav[j]) else 'R02',solver.value(relay_start[j]),solver.value(relay_end[j]),f'R-B{j+1:02d}'])
print('wrote luna_flex_transport.csv and luna_flex_relays.csv',flush=True)


