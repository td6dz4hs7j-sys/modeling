"""Render Q2 figures from the selected balanced schedule and comparison CSVs."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import Patch
import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "figures/q2_strategy_scenarios"
OUT.mkdir(parents=True,exist_ok=True)
sources = {}


def source(rel):
    p = ROOT / rel
    sources[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
    return p


def save(fig, name):
    fig.savefig(OUT / f"{name}.png", dpi=360, bbox_inches="tight")
    fig.savefig(OUT / f"{name}.svg", bbox_inches="tight")
    plt.close(fig)


plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9,
                     "axes.spines.top": False, "axes.spines.right": False})
blue, orange, green, purple = "#0072B2", "#D55E00", "#009E73", "#CC79A7"
wb = load_workbook(source("results/q2_strategy_scenarios/Q2_综合均衡最终结果.xlsx"),
                   read_only=True, data_only=True)
sorties = [r for r in wb["Q2_运输架次"].values if isinstance(r[0], str) and r[0].startswith("Q2-")]
boxes = [r for r in wb["Q2_逐箱交付"].values if isinstance(r[0], str) and r[0].startswith("S")]
assert len(sorties) == 23 and len(boxes) == 80
assert len({r[0] for r in boxes}) == 80
assert abs(sum(float(r[7]) for r in sorties) - 65.8690914160528) < 1e-7
assert abs(max(float(r[6]) for r in sorties) / 60 - 97.7058717164508) < 1e-7

# Schedule: the selected 23 sorties, colored by aircraft type.
fig, ax = plt.subplots(figsize=(7.1, 5.4), layout="constrained")
color = {"A": blue, "B": orange, "C": green}
for i, row in enumerate(sorties):
    ax.barh(i, (float(row[6]) - float(row[4])) / 60, left=float(row[4]) / 60,
            height=.72, color=color[str(row[2])])
ax.set(yticks=np.arange(23), yticklabels=[r[0] for r in sorties],
       xlabel="Time from start (min)", title="Selected balanced schedule: 23 sorties")
ax.invert_yaxis()
ax.grid(axis="x", alpha=.2)
ax.legend(handles=[Patch(facecolor=color[typ], label=f"Type {typ}") for typ in "ABC"],
          loc="upper center", bbox_to_anchor=(.5, -.10), ncol=3, frameon=False)
save(fig, "q2_balanced_schedule")

# Demand-linked box slack; sorted values avoid hiding the minimum.
raw = load_workbook(source("input/数据/无人机应急物资运输基础数据/物资需求与配送时限.xlsx"),
                    read_only=True, data_only=True)
lookup = {r[0]: r for r in raw["逐箱货箱清单"].values
          if isinstance(r[0], str) and r[0].startswith("S")}
assert {r[0] for r in boxes} == set(lookup)
slack = np.sort([(float(lookup[r[0]][7]) - float(r[3])) / 60 for r in boxes])
assert abs(slack[0] - 1.381755287271896) < 1e-7
fig, ax = plt.subplots(figsize=(7.1, 3.0), layout="constrained")
ax.plot(np.arange(1, 81), slack, color=blue, lw=1.8)
ax.axhline(0, color="black", lw=.8)
ax.scatter([1], [slack[0]], color=orange, s=45, zorder=3)
ax.annotate(f"minimum {slack[0]:.2f} min", (1, slack[0]), xytext=(9, 9),
            textcoords="offset points", fontsize=8)
ax.set(xlabel="Boxes ranked by expected-time slack", ylabel="Slack (min)",
       title="All 80 boxes meet their expected time", xlim=(1, 80))
ax.grid(alpha=.2)
save(fig, "q2_balanced_delivery_slack")

# Resource occupancy: flight and two-stage charging for each physical battery.
capacity = {"A": 4.5, "B": 4.0, "C": 8.0}
full_charge = {"A": 1800, "B": 2400, "C": 3000}
ids = sorted({str(r[3]) for r in sorties})
fig, ax = plt.subplots(figsize=(7.1, 4.0), layout="constrained")
for row in sorties:
    typ, battery = str(row[2]), str(row[3])
    start, ret, energy = map(float, (row[4], row[6], row[7]))
    soc = 1 - energy / capacity[typ]
    full = full_charge[typ]
    charge = (full * (.65 * (.9 - soc) / .9 + .35) if soc < .9
              else full * .35 * (1 - soc) / .1)
    y = ids.index(battery)
    ax.barh(y, (ret - start) / 3600, left=start / 3600, height=.7, color=color[typ])
    ax.barh(y, charge / 3600, left=ret / 3600, height=.7, color=color[typ], alpha=.25)
ax.set(yticks=np.arange(len(ids)), yticklabels=ids, xlabel="Time from start (h)",
       title="Selected schedule: battery flight and recharge")
ax.invert_yaxis()
ax.grid(axis="x", alpha=.2)
ax.text(.99, .02, "solid: flight   pale: recharge", transform=ax.transAxes,
        ha="right", fontsize=8)
save(fig, "q2_balanced_battery_occupancy")

# Same algorithm, seed pool and budget: four weight vectors on the same instance.
with source("results/q2_strategy_scenarios/weighted_scheme_comparison.csv").open(
        encoding="utf-8-sig", newline="") as f:
    weighted = list(csv.DictReader(f))
assert len(weighted) == 4 and all(r["validator"] == "PASS" and int(r["boxes_on_time"]) == 80 for r in weighted)
labels = {"架次优先": "Sortie priority", "时间优先": "Time priority",
          "能耗优先": "Energy priority", "综合均衡": "Balanced (selected)"}
fig, ax = plt.subplots(figsize=(7.1, 3.8), layout="constrained")
for r in weighted:
    selected = r["scheme"] == "综合均衡"
    energy, minute = float(r["energy_kWh"]), float(r["Cmax_min"])
    ax.scatter(energy, minute, s=180 if selected else 90,
               color=orange if selected else blue, marker="*" if selected else "o",
               edgecolor="black", linewidth=.5, zorder=3,
               label=labels[r["scheme"]] + f"  N={r['sorties']}")
ax.set(xlabel="Transport energy (kWh)", ylabel="Latest return (min)",
       title="Strict PASS solutions under four objective weights",
       xlim=(65.55, 68.55), ylim=(95.25, 101.55))
ax.legend(loc="upper right", fontsize=8, frameon=True)
ax.grid(alpha=.2)
save(fig, "q2_weighted_tradeoff")

# Distinguish algorithm/constructive routes from the controlled weight experiment.
with source("results/q2_strategy_scenarios/algorithm_strategy_comparison.csv").open(
        encoding="utf-8-sig", newline="") as f:
    strategies = list(csv.DictReader(f))
assert len(strategies) == 6 and all(r["validator"] == "PASS" for r in strategies)
fig, ax = plt.subplots(figsize=(7.1, 3.8), layout="constrained")
short = ["EDD", "22-sortie", "Route split", "Local time", "Energy refine", "Weighted search"]
for i, r in enumerate(strategies):
    ax.scatter(float(r["energy_kWh"]), float(r["Cmax_min"]),
               s=65 + (20 if r["sorties"] == "22" else 0),
               marker="s" if r["sorties"] == "22" else "o",
               color=purple if i == 5 else green, edgecolor="black", linewidth=.4)
    ax.annotate(short[i], (float(r["energy_kWh"]), float(r["Cmax_min"])),
                xytext=(5, 5), textcoords="offset points", fontsize=8)
ax.set(xlabel="Transport energy (kWh)", ylabel="Latest return (min)",
       title="Verified results from different construction strategies",
       xlim=(65.65, 69.65))
ax.grid(alpha=.2)
save(fig, "q2_strategy_comparison")

(Path(__file__).resolve().parent / "q2_balanced_figure_sources.json").write_text(
    json.dumps({"input_sha256": sources, "selected_sorties": len(sorties),
                "selected_boxes": len(boxes), "min_box_slack_min": float(slack[0])},
               ensure_ascii=False, indent=2), encoding="utf-8")
print("Q2 balanced figures: 5 PNG + 5 SVG; 23 sorties, 80 boxes")
