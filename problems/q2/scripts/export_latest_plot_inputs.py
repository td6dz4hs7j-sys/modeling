"""Export the frozen balanced MAT timeline and node coordinates for plotting."""
import argparse,csv,hashlib,json
from pathlib import Path
import h5py
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--source-root',type=Path,default=ROOT);args=ap.parse_args()
    out=ROOT/'results/q2_strategy_scenarios';mat=out/'balanced_best_pass.mat'
    nodes=args.source_root/'input/数据/无人机应急物资运输基础数据/调度中心与服务区.xlsx'
    w=load_workbook(nodes,data_only=True);raw=list(w['数据'].values)
    nrows=[['index','id','lon','lat','altitude_m']]+[[i+1,*[r[j] for j in (0,2,3,4)]] for i,r in enumerate([raw[2],*raw[6:21]])]
    timeline=[['sortie_id','uav_id','battery_id','type','nodes','boxes','prep_start_s','prep_end_s','load_start_s','load_end_s','takeoff_s','return_s','battery_ready_s']]
    official=load_workbook(out/'Q2_综合均衡最终结果.xlsx',data_only=True)
    listed=[r for r in official['Q2_运输架次'].values if isinstance(r[0],str) and r[0].startswith('Q2-')]
    assert len(listed)==23
    with h5py.File(mat) as f:
        routes=f['best/routes']
        def data(key,i):return f[routes[key][0,i]][()]
        def number(key,i):return float(data(key,i).item())
        for i in range(routes['typeIdx'].shape[1]):
            boxes=data('boxIdx',i).size;prep=number('prepStart_s',i);load=number('loadStart_s',i)
            row=listed[i]
            assert row[0]==f'Q2-{i+1:03d}' and abs(float(row[4])-number('takeoff_s',i))<1e-5
            assert abs(float(row[6])-number('return_s',i))<1e-5
            timeline.append([row[0],row[1],row[3],'ABC'[int(number('typeIdx',i))-1],
                             ';'.join(str(int(n)) for n in data('nodeOrder',i).flatten()),boxes,prep,prep+300,load,load+30*boxes,
                             number('takeoff_s',i),number('return_s',i),number('batteryReady_s',i)])
    assert len(timeline)==24 and len(nrows)==17
    assert abs(max(r[-2] for r in timeline[1:])/60-97.7058717164508)<1e-7
    for r in timeline[1:]:assert max(r[7],r[9])<=r[10]+1e-5
    for name,rows in [('latest_route_timeline.csv',timeline),('latest_nodes.csv',nrows)]:
        with (out/name).open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
    d={'balanced_mat_sha256':hashlib.sha256(mat.read_bytes()).hexdigest(),
       'official_workbook_sha256':hashlib.sha256((out/'Q2_综合均衡最终结果.xlsx').read_bytes()).hexdigest(),
       'source_node_workbook_sha256':hashlib.sha256(nodes.read_bytes()).hexdigest(),
       'plot_csv_sha256':{n:hashlib.sha256((out/n).read_bytes()).hexdigest() for n in ('latest_route_timeline.csv','latest_nodes.csv')},
       'preparation_s':300,'load_s_per_box':30,'note':'All three types use these preparation/loading values in the original attachment.'}
    (out/'latest_plot_input_provenance.json').write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
    print('PASS: exported 23 frozen routes and 16 nodes for current-result figures')
if __name__=='__main__':main()
