"""Detailed, reproducible multi-objective analysis of strict-PASS Q2 schedules."""
from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "results" / "q2_99_review"


def read(name: str) -> list[dict[str, str]]:
    with (OUT / name).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def write(name: str, headers: list[str], data: list[list[object]]) -> None:
    with (OUT / name).open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(data)


def f(item: dict, key: str) -> float:
    return float(item[key])


def peak(intervals: list[tuple[float, float]]) -> int:
    events = []
    for start, end in intervals:
        if end > start:
            events.extend([(start, 1), (end, -1)])
    active = maximum = 0
    for _, delta in sorted(events, key=lambda event: (event[0], event[1])):
        active += delta
        maximum = max(maximum, active)
    return maximum


def main() -> None:
    report = (OUT / "audit_report.txt").read_text(encoding="utf-8-sig")
    assert "status=PASS" in report and "boxes_on_time=80/80" in report
    methods = read("algorithm_comparison.csv")
    assert all(item["validator"] == "PASS" for item in methods)
    sorties = read("sorties_23.csv")
    boxes = read("boxes_80.csv")
    batteries = read("battery_timeline.csv")
    areas = read("area_summary.csv")
    assert len(sorties) == 23 and len(boxes) == 80 and len(areas) == 15
    final = methods[-1]
    cmax = f(final, "Cmax_min")
    energy = f(final, "energy_kWh")
    assert abs(max(f(s, "return_min") for s in sorties) - cmax) < 1e-7
    assert abs(sum(f(s, "energy_kWh") for s in sorties) - energy) < 1e-7
    assert all(f(b, "expected_margin_min") >= -1e-7 for b in boxes)

    # Hard deadline and N<=23 are admission rules. Weighted objective applies
    # only after a schedule passes those rules.
    goal_rows = []
    for item in methods:
        time_gap = max(0, (f(item, "Cmax_min") - 94) / 94) * 100
        energy_gap = max(0, (f(item, "energy_kWh") - 67) / 67) * 100
        goal_rows.append([item["method"], f(item, "Cmax_min"), f(item, "energy_kWh"),
                          time_gap, energy_gap, f(item, "Cmax_min") - 94,
                          f(item, "energy_kWh") - 67])
    write("goal_gaps.csv", ["method", "Cmax_min", "energy_kWh", "time_gap_pct_vs_94",
                            "energy_gap_pct_vs_67", "time_gap_min", "energy_gap_kWh"], goal_rows)

    target_rows = []
    for a in (0, .25, .5, .75, 1):
        scores = [(item["method"], a * row[3] + (1 - a) * row[4])
                  for item, row in zip(methods, goal_rows)]
        winner = min(scores, key=lambda pair: pair[1])[0]
        for method, score in scores:
            target_rows.append([a, 1 - a, method, score, method == winner])
    write("weight_time_energy.csv", ["time_weight", "energy_weight", "method",
                                     "weighted_target_gap_pct", "best_under_weights"], target_rows)

    # If sortie count itself is a soft objective, normalize each metric over
    # the four strict-PASS alternatives. Only 22-sortie and final remain Pareto.
    mins = {key: min(f(x, key) for x in methods) for key in ("sorties", "Cmax_min", "energy_kWh")}
    maxs = {key: max(f(x, key) for x in methods) for key in mins}
    norm = {}
    for item in methods:
        norm[item["method"]] = {key: (f(item, key) - mins[key]) / (maxs[key] - mins[key]) for key in mins}
    preferences = [
        ("time_priority", .10, .65, .25),
        ("energy_priority", .10, .20, .70),
        ("balanced", 1 / 3, 1 / 3, 1 / 3),
        ("sortie_priority", .60, .20, .20),
        ("strong_sortie_priority", .75, .15, .10),
    ]
    scenario_rows = []
    for name, wn, wt, we in preferences:
        scored = [(item["method"], wn * norm[item["method"]]["sorties"]
                   + wt * norm[item["method"]]["Cmax_min"]
                   + we * norm[item["method"]]["energy_kWh"]) for item in methods]
        best_name = min(scored, key=lambda pair: pair[1])[0]
        for method, score in scored:
            scenario_rows.append([name, wn, wt, we, method, score, method == best_name])
    write("weight_three_objectives.csv", ["scenario", "sortie_weight", "time_weight",
                                           "energy_weight", "method", "normalized_loss",
                                           "best_under_weights"], scenario_rows)
    norm_rows = [[m, n["sorties"], n["Cmax_min"], n["energy_kWh"]] for m, n in norm.items()]
    write("objective_normalization.csv", ["method", "normalized_sorties", "normalized_Cmax",
                                           "normalized_energy"], norm_rows)

    # UAV and UAV-type workload, charging and deadline evidence.
    by_type = defaultdict(list)
    by_uav = defaultdict(list)
    for item in sorties:
        by_type[item["uav_type"]].append(item)
        by_uav[item["uav_id"]].append(item)
    type_rows = []
    for typ in sorted(by_type):
        group = by_type[typ]
        type_rows.append([typ, len({x["uav_id"] for x in group}), len(group),
                          sum(int(x["box_count"]) for x in group),
                          sum(int(x["service_area_count"]) > 1 for x in group),
                          sum(f(x, "energy_kWh") for x in group),
                          sum(f(x, "operation_s") for x in group) / 60,
                          max(f(x, "return_min") for x in group)])
    write("type_summary.csv", ["type", "physical_uavs", "sorties", "boxes", "multi_area_sorties",
                               "energy_kWh", "total_flight_handover_min", "last_return_min"], type_rows)
    overlap_rows = []
    for uav, group in sorted(by_uav.items()):
        group.sort(key=lambda item: f(item, "takeoff_min"))
        for prev, current in zip(group, group[1:]):
            overlap = max(0, min(f(current, "prep_end_s"), f(prev, "return_s"))
                          - max(f(current, "prep_start_s"), f(prev, "takeoff_s"))) / 60
            assert f(current, "prep_start_s") >= f(prev, "takeoff_s") - 1e-6
            assert f(current, "load_start_s") >= f(prev, "return_s") - 1e-6
            overlap_rows.append([uav, prev["sortie_id"], current["sortie_id"], overlap,
                                 f(current, "load_start_min") - f(prev, "return_min")])
    write("preparation_overlap.csv", ["uav_id", "previous_sortie", "next_sortie",
                                      "prep_overlaps_prior_flight_min", "load_after_return_gap_min"], overlap_rows)
    categories = defaultdict(list)
    for item in boxes:
        if item["category"] == "医疗物资":
            cat = "medical"
        elif item["hard_deadline_s"] not in ("Inf", "inf", ""):
            cat = "first_batch_nonmedical"
        else:
            cat = "other"
        categories[cat].append(item)
    deadline_rows = []
    for cat in ("medical", "first_batch_nonmedical", "other"):
        group = categories[cat]
        hard = [f(x, "hard_margin_min") for x in group if x["hard_margin_min"] not in ("Inf", "inf", "")]
        deadline_rows.append([cat, len(group), sum(f(x, "expected_margin_min") >= -1e-7 for x in group),
                              min(f(x, "expected_margin_min") for x in group),
                              min(hard) if hard else "", mean(f(x, "arrival_min") for x in group)])
    write("deadline_category_summary.csv", ["category", "boxes", "expected_on_time",
                                            "min_expected_margin_min", "min_hard_margin_min",
                                            "mean_arrival_min"], deadline_rows)
    flying = [(f(x, "takeoff_min"), f(x, "return_min")) for x in sorties]
    charging = [(f(x, "return_min"), f(x, "recharged_min")) for x in batteries
                if x["sortie_id"] and x["recharged_min"]]
    busy_battery = [(f(x, "takeoff_min"), f(x, "recharged_min")) for x in batteries
                    if x["sortie_id"] and x["recharged_min"]]
    resource_rows = [
        ["peak_flying_uavs", peak(flying), "count"],
        ["peak_charging_batteries", peak(charging), "count"],
        ["peak_occupied_batteries", peak(busy_battery), "count"],
        ["distinct_used_batteries", len({x["battery_id"] for x in batteries if x["sortie_id"]}), "count"],
        ["consecutive_prep_overlap_total", sum(row[3] for row in overlap_rows), "min; not additive Cmax saving"],
        ["consecutive_prep_overlap_count", sum(row[3] > 1e-8 for row in overlap_rows), "count"],
        ["minimum_expected_margin", min(f(x, "expected_margin_min") for x in boxes), "min"],
        ["minimum_energy_reserve_margin", min(f(x, "reserve_margin_kWh") for x in sorties), "kWh"],
    ]
    write("resource_and_margin_summary.csv", ["metric", "value", "unit_or_note"], resource_rows)

    old_path = ROOT / "results" / "q2_prestage_23_refine" / "Q2_99分钟_原版结果表.xlsx"
    old_book = load_workbook(old_path, read_only=True, data_only=True)
    old_sheet = old_book["23架次排程"]
    headings = [cell.value for cell in old_sheet[1]]
    old = {str(row[0]): dict(zip(headings, row)) for row in old_sheet.iter_rows(min_row=2, values_only=True)}
    delta_rows = []
    for item in sorties:
        previous = old[item["sortie_id"]]
        delta_rows.append([item["sortie_id"], previous["uav_id"], item["uav_id"],
                           previous["service_order"], item["service_order"],
                           previous["box_ids"] != item["box_ids"],
                           f(item, "return_min") - float(previous["return_min"]),
                           f(item, "energy_kWh") - float(previous["energy_kWh"])])
    write("sortie_old_new_delta.csv", ["sortie_id", "old_uav", "new_uav", "old_service_order",
                                      "new_service_order", "box_assignment_changed",
                                      "return_change_min", "energy_change_kWh"], delta_rows)

    book = load_workbook(OUT / "Q2_最终核验与结果表.xlsx")
    extras = [
        ("目标差距", "goal_gaps.csv"), ("时间能耗权重", "weight_time_energy.csv"),
        ("三目标权重", "weight_three_objectives.csv"), ("归一化口径", "objective_normalization.csv"),
        ("机型汇总", "type_summary.csv"), ("准备重叠", "preparation_overlap.csv"),
        ("时限类别", "deadline_category_summary.csv"), ("资源峰值", "resource_and_margin_summary.csv"),
        ("旧新架次差异", "sortie_old_new_delta.csv"),
    ]
    for title, name in extras:
        ws = book.create_sheet(title)
        with (OUT / name).open(encoding="utf-8-sig", newline="") as fh:
            for row in csv.reader(fh):
                converted = []
                for value in row:
                    try:
                        converted.append(float(value) if value and value[0] in "-+.0123456789" else value)
                    except ValueError:
                        converted.append(value)
                ws.append(converted)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(
                45, max(13, max(len(str(c.value or "")) for c in col[: min(len(col), 80)]) + 2)
            )
    result_path = OUT / "Q2_多目标综合对比.xlsx"
    book.save(result_path)
    check = load_workbook(result_path, read_only=True, data_only=True)
    assert check["目标差距"].max_row == len(methods) + 1
    assert check["准备重叠"].max_row == len(overlap_rows) + 1
    assert check["旧新架次差异"].max_row == 24

    e22 = norm[methods[0]["method"]]["energy_kWh"]
    scenario_names = {"time_priority": "时间优先", "energy_priority": "能耗优先",
                      "balanced": "三目标均衡", "sortie_priority": "架次优先",
                      "strong_sortie_priority": "强架次优先"}
    method_names = {"22-sortie time-first search": "22 架次时间优先搜索",
                    "23-sortie energy refinement": "23 架次最终节能方案"}
    best_scenarios = []
    for name, wn, wt, we in preferences:
        winner = next(row[4] for row in scenario_rows if row[0] == name and row[6])
        best_scenarios.append(f"| {scenario_names[name]} | {wn:.3f} | {wt:.3f} | {we:.3f} | {method_names.get(winner,winner)} |")
    type_lines = [f"| {r[0]} | {r[1]} | {r[2]} | {r[3]} | {r[5]:.4f} | {r[7]:.4f} |" for r in type_rows]
    category_lines = [f"| {r[0]} | {r[1]} | {r[2]} | {r[3]:.4f} | {r[4] if r[4] == '' else f'{r[4]:.4f}'} |" for r in deadline_rows]
    note = f"""# 问题二：多目标权重与方案对比

## 口径与结论

四个方案均由 MATLAB 严格 validator 再验为 PASS，80/80 箱按期。最终方案为 **23 架次、{cmax:.6f} min、{energy:.6f} kWh**。相对原 99.0431 min 方案，时间缩短 {f(methods[-2], 'Cmax_min')-cmax:.4f} min、能耗下降 {f(methods[-2], 'energy_kWh')-energy:.4f} kWh，且未增加架次。原版逐架次数据与最终数据的差异见 `sortie_old_new_delta.csv`。

先设硬约束：80 箱按期、全部硬时限和资源安全约束成立、架次数不超过 23。随后用时间和能耗权重选择方案。最终方案在两个指标上均优于其他三个 PASS 方案，所以在此口径下，不论时间/能耗权重如何取非负且和为 1，最终方案都排第一。这个结论仅针对当前四个已验证候选，不是全局最优证明。

## 与目标指标的距离

将 94 min、67 kWh 作为外部参考目标，不把它当作已验证可行解。最终方案仍差 **{cmax-94:.4f} min** 和 **{energy-67:.4f} kWh**。归一化目标差距定义为 `max(0,(Cmax-94)/94)` 与 `max(0,(E-67)/67)`；逐方案数值和 0%、25%、50%、75%、100% 时间权重的分数见 `goal_gaps.csv`、`weight_time_energy.csv`。

## 把架次数作为软目标时

如果愿意用多 1 架次换时间和能耗，就把 N、Cmax、E 各自在四个 PASS 方案的最小值和最大值之间归一化为 0 到 1，然后最小化 `wN*N' + wT*Cmax' + wE*E'`，权重非负且和为 1。最终 23 架次方案与 22 架次方案是当前候选集中的两个非支配方案；另外两个 23 架次方案均被最终方案支配。

22 架次方案的归一化时间损失为 1，能耗损失为 **{e22:.4f}**；最终方案的架次损失为 1，时间及能耗损失为 0。因此最终方案胜出的边界为 **`wN < wT + {e22:.4f}·wE`**，相等时两者打平。该边界依赖当前候选范围的归一化，若加入新方案需重算。

| 权重情景 | wN | wT | wE | 首选方案 |
|---|---:|---:|---:|---|
{chr(10).join(best_scenarios)}

完整四方案评分见 `weight_three_objectives.csv`，权重分界图见 `figures/weight_balance_region.png`（附 SVG 及灰度图）。这些情景是明确的决策偏好示例，不是题目给定权重。

## 资源与交付结构

| 机型 | 实体机 | 架次 | 货箱 | 能耗/kWh | 最晚返航/min |
|---|---:|---:|---:|---:|---:|
{chr(10).join(type_lines)}

| 时限类别 | 货箱 | 期望按期 | 最小期望裕度/min | 最小硬时限裕度/min |
|---|---:|---:|---:|---:|
{chr(10).join(category_lines)}

最大同时飞行 {peak(flying)} 架，最大同时充电 {peak(charging)} 组电池，实际使用 {len({x['battery_id'] for x in batteries if x['sortie_id']})} 组不同电池。上一架次飞行与下一架次固定准备的重叠共 **{sum(row[3] for row in overlap_rows):.4f} min**，这是各机各次的重叠量之和，不能解释为 Cmax 同幅缩短。最紧的期望交付时限裕度为 **{min(f(x, 'expected_margin_min') for x in boxes):.4f} min**；最小返航安全能量裕度为 **{min(f(x, 'reserve_margin_kWh') for x in sorties):.4f} kWh**。

按服务区的逐箱交付、最晚到达和最小裕度在 `area_summary.csv`；8 架无人机的工作量在 `uav_summary.csv`；14 组共享电池各次起飞、返航、充电完成及 SOC 在 `battery_timeline.csv`。图见 `figures/`。固定准备在上一架起飞后提前进行是扩展排程假设，相关重叠时段见 `preparation_overlap.csv` 和无人机甘特图。未提供工位数量，本方案按可并行准备核算。

## 可复核文件

`Q2_多目标综合对比.xlsx` 包含原有逐架次、逐箱、算法表及本次目标差距、权重敏感性、机型、时限、资源峰值和新旧逐架次差异。全部工作表均另有 CSV；冻结排程在 `final_best_pass.mat`，独立核验与破坏测试记录在 `audit_report.txt` 和 `validator_mutations.csv`。
"""
    (OUT / "多目标权重与方案对比.md").write_text(note, encoding="utf-8")
    print(result_path)


if __name__ == "__main__":
    main()
