"""Package MATLAB-validated Q2 tables without recalculating optimization results."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "results" / "q2_99_review"

TABLES = [
    ("算法对比", "algorithm_comparison.csv"),
    ("同路线排序对照", "same_routes_order_baselines.csv"),
    ("23架次排程", "sorties_23.csv"),
    ("80箱交付", "boxes_80.csv"),
    ("无人机汇总", "uav_summary.csv"),
    ("共享电池时序", "battery_timeline.csv"),
    ("服务区汇总", "area_summary.csv"),
    ("校验破坏测试", "validator_mutations.csv"),
]


def rows(name: str) -> list[list[str]]:
    with (OUT / name).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


def typed(value: str):
    if value == "":
        return None
    try:
        if value.isdigit() and not (len(value) > 1 and value[0] == "0"):
            return int(value)
        if any(c in value for c in ".eE") and not any(c.isalpha() for c in value.replace("e", "").replace("E", "")):
            return float(value)
    except ValueError:
        pass
    return value


def main() -> None:
    assert (OUT / "run_status.txt").read_text().strip() == "0"
    assert "status=PASS" in (OUT / "audit_report.txt").read_text(encoding="utf-8")
    with (OUT / "algorithm_comparison.csv").open(encoding="utf-8-sig", newline="") as fh:
        methods = list(csv.DictReader(fh))
    with (OUT / "same_routes_order_baselines.csv").open(encoding="utf-8-sig", newline="") as fh:
        order_cases = list(csv.DictReader(fh))
    with (OUT / "sorties_23.csv").open(encoding="utf-8-sig", newline="") as fh:
        sorties = list(csv.DictReader(fh))
    with (OUT / "boxes_80.csv").open(encoding="utf-8-sig", newline="") as fh:
        boxes = list(csv.DictReader(fh))
    final = methods[-1]
    previous = methods[-2]
    cmax = float(final["Cmax_min"])
    energy = float(final["energy_kWh"])
    assert final["validator"] == "PASS" and int(final["boxes_on_time"]) == 80
    assert cmax <= 99.043089527783 + 1e-8
    assert energy <= 69.251959983412 + 1e-8
    wb = Workbook()
    wb.remove(wb.active)
    for title, filename in TABLES:
        data = rows(filename)
        ws = wb.create_sheet(title)
        for line in data:
            ws.append([typed(v) for v in line])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 34
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            width = min(48, max(13, max(len(str(c.value or "")) for c in col[: min(len(col), 82)]) + 2))
            ws.column_dimensions[letter].width = width
            for cell in col[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
    notes = wb.create_sheet("口径说明", 0)
    for item in [
        ("指标", "数值或说明"),
        ("严格校验", "PASS；80/80 箱按期；23 架次"),
        ("最晚返航", f"{cmax:.12f} min"),
        ("总能耗", f"{energy:.12f} kWh"),
        ("排程假设", "下一架固定准备可在上一架起飞后开始；装载须等返航；充电可与准备并行，起飞前电池充满。"),
        ("算法比较", "已存档候选方案用同一输入和相同排程口径重新执行严格 validator。运行预算不同，不宣称公平竞速或全局最优。"),
        ("排序消融", "固定最终23条路线、货箱和机型，只改路线排序。不可行结果仅供诊断，不算可行成果。"),
        ("时间定义", "最晚返航 Cmax；CSV 时间字段带 _s 为秒，带 _min 为分钟。"),
        ("外部目标", "23 架次、94 min、67 kWh 仅有指标，未获得排程文件，未纳入已验证算法比较。"),
    ]:
        notes.append(item)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 105
    notes.freeze_panes = "A2"
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor="16324F")
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    path = OUT / "Q2_最终核验与结果表.xlsx"
    wb.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    assert check["算法对比"].max_row == len(methods) + 1
    assert check["同路线排序对照"].max_row == 5
    assert check["23架次排程"].max_row == 24
    assert check["80箱交付"].max_row == 81
    assert check["共享电池时序"].max_row >= 15

    names = {
        "22-sortie time-first search": "22 架次时间优先搜索",
        "23-sortie route split": "23 架次路线拆分",
        "23-sortie local refinement": "23 架次局部精修",
        "23-sortie energy refinement": "23 架次能耗精修（最终）",
    }
    method_lines = [
        "| 方案 | 架次 | 最晚返航/min | 能耗/kWh | 按期交付 | 多区架次 | 严格校验 |",
        "|---|---:|---:|---:|---:|---:|---|",
    ]
    for item in methods:
        method_lines.append(
            f"| {names.get(item['method'], item['method'])} | {item['sorties']} | "
            f"{float(item['Cmax_min']):.4f} | {float(item['energy_kWh']):.4f} | "
            f"{item['boxes_on_time']}/80 | {item['multi_area_sorties']} | {item['validator']} |"
        )
    order_names = {
        "earliest_due_date": "最早期限优先",
        "shortest_route_first": "短路线优先",
        "input_box_order": "输入货箱顺序",
        "optimized_route_order": "优化路线顺序",
    }
    order_lines = [
        "| 路线排序规则 | 最晚返航/min | 按期货箱 | 严格校验 |",
        "|---|---:|---:|---|",
    ]
    for item in order_cases:
        state = "PASS" if item["strict_PASS"] == "1" else "FAIL"
        order_lines.append(
            f"| {order_names.get(item['route_order_rule'], item['route_order_rule'])} | "
            f"{float(item['Cmax_min']):.4f} | {item['boxes_on_time']}/80 | {state} |"
        )
    last = max(sorties, key=lambda item: float(item["return_min"]))
    critical = sorted(
        (item for item in sorties if item["uav_id"] == last["uav_id"]),
        key=lambda item: float(item["takeoff_min"]),
    )
    chain = "、".join(f"{item['sortie_id']}（返航 {float(item['return_min']):.4f} min）" for item in critical)
    min_due = min(float(item["expected_margin_min"]) for item in boxes)
    min_reserve = min(float(item["reserve_margin_kWh"]) for item in sorties)
    delta_t = float(previous["Cmax_min"]) - cmax
    delta_e = float(previous["energy_kWh"]) - energy
    summary = f"""# 问题二：最终方案核验与图表说明

## 核验结论

从原始工作簿读取 15 个服务区、80 个不可拆货箱、8 架实体无人机和 14 组共享电池，并重新执行独立路线能耗、时间和资源核算。严格 validator **PASS**：23 架次，最晚返航 **{cmax:.8f} min**，总能耗 **{energy:.8f} kWh**，80/80 箱按期，{final['multi_area_sorties']} 架次访问多个服务区。最小期望时限裕度 {min_due:.4f} min，最小返航安全能量裕度 {min_reserve:.4f} kWh。固定准备、装载、充电、实体无人机及电池时序均通过检查，四个针对性破坏样例被 validator 拒绝。

**排程口径**：下一架次固定准备可以在该无人机上一架次起飞后开始；装载必须等其返航；准备、装载与共享电池充电可并行，起飞前均须完成。提前准备是本次探索的扩展排程假设。最晚返航是各架次返航时刻的最大值。

## 同一排程口径下的算法方案对比

{chr(10).join(method_lines)}

能耗精修相对上一个 99.0431 分钟 PASS 方案，最晚返航提前 **{delta_t:.4f} min**，能耗减少 **{delta_e:.4f} kWh**，架次和按期箱数不变。与 22 架次方案相比，本方案更快，但架次更多；与路线拆分方案相比，本方案时间与能耗都更低。比较对象均重新通过严格 validator，但搜索预算不同，不能由此宣称全局最优或公平竞速。外部的 23 架次、94 min、67 kWh 只有指标、没有可复核排程，不纳入 PASS 对比。

### 固定最终路线的排序消融

固定最终 23 条路线、货箱和机型，只改变路线输入次序，资源排程器相同；因此各行能耗均为 {energy:.4f} kWh。

{chr(10).join(order_lines)}

FAIL 行仅作诊断，时间不作为可行成绩；具体违反约束见 `same_routes_order_baselines.csv`。

## 关键完成链

最后返航的实体无人机为 {last['uav_id']}：{chain}。后继固定准备可与前次飞行重叠；装载须从前次返航后开始。见 `sorties_23.csv` 中的 `prep_start_s`、`load_start_s`、`takeoff_s`、`return_s`。

## 图与表

图由 MATLAB 直接依据最终冻结方案生成，均有 SVG、300 dpi PNG 和灰度质检图：

- `figures/uav_prep_load_flight.png`：准备、装载、飞行与交付的重叠时序。
- `figures/battery_flight_charge.png`：14 组共享电池的任务占用和返航后充电。
- `figures/routes_geographic.png`：23 架次实际地理路线。
- `figures/box_deadline_margins.png`：80 箱期望时限裕度。
- `figures/algorithm_comparison.png`：已验证方案的时间、架次与能耗对比。
- `figures/sortie_energy_reserve.png`：逐架次能耗及安全余量。
- `figures/area_demand.png`、`figures/box_mass_distribution.png`：需求结构。
- `figures/route_stops_distribution.png`：路线服务区数量结构。

`Q2_最终核验与结果表.xlsx` 包含算法比较、排序消融、逐架次、逐箱、无人机、电池、服务区和校验测试工作表。各表另有 CSV；`audit_report.txt`、`validator_mutations.csv`、`matlab.log` 留存核验记录。原 99.0431 分钟方案的冻结 MAT 文件仍在 `results/q2_prestage_23_refine/best_pass.mat`。
"""
    (OUT / "结果说明与图注.md").write_text(summary, encoding="utf-8")
    print(path)


if __name__ == "__main__":
    main()
