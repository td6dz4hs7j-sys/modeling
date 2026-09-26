"""Reference-style tables for independently MATLAB-validated Q2 strategies."""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "results" / "q2_strategy_scenarios"
FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")


def read(name: str) -> list[dict[str, str]]:
    with (OUT / name).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def display_table(name: str, title: str, subtitle: str, headers: list[str],
                  rows: list[list[str]], widths: list[int], note: str,
                  highlight_row: int | None = None) -> None:
    margin = 54
    table_width = sum(widths)
    width = table_width + 2 * margin
    row_height = 108
    header_height = 86
    height = 78 + 54 + header_height + len(rows) * row_height + 118
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(str(FONT_BOLD), 34)
    header_font = ImageFont.truetype(str(FONT_BOLD), 24)
    body_font = ImageFont.truetype(str(FONT), 22)
    body_bold = ImageFont.truetype(str(FONT_BOLD), 22)
    small_font = ImageFont.truetype(str(FONT), 18)
    draw.text((margin, 30), title, font=title_font, fill="#15283D")
    draw.text((margin, 83), subtitle, font=small_font, fill="#596575")
    y = 132
    draw.rounded_rectangle((margin, y, margin + table_width, y + header_height),
                           radius=9, fill="#EAF0F6")
    x = margin
    for text, col_width in zip(headers, widths):
        wrapped = wrap(draw, text, header_font, col_width - 18)
        draw.multiline_text((x + 11, y + 18), wrapped, font=header_font,
                            fill="#20384F", spacing=3)
        x += col_width
    y += header_height
    for r, line in enumerate(rows):
        fill = "#EAF6F0" if r == highlight_row else ("#FFFFFF" if r % 2 == 0 else "#F8FAFC")
        draw.rectangle((margin, y, margin + table_width, y + row_height), fill=fill)
        x = margin
        for c, (text, col_width) in enumerate(zip(line, widths)):
            font = body_bold if c == 0 or r == highlight_row and c in (3, 4, 5) else body_font
            wrapped = wrap(draw, str(text), font, col_width - 18)
            draw.multiline_text((x + 11, y + 18), wrapped, font=font,
                                fill="#133047" if r != highlight_row else "#154F38", spacing=5)
            x += col_width
        draw.line((margin, y + row_height, margin + table_width, y + row_height),
                  fill="#DDE3EA", width=2)
        y += row_height
    draw.text((margin, y + 23), wrap(draw, note, small_font, table_width - 8),
              font=small_font, fill="#5D6771", spacing=5)
    img.save(OUT / name, dpi=(300, 300), optimize=True)


def wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont,
         max_width: int) -> str:
    parts = []
    for segment in text.split("\n"):
        current = ""
        for char in segment:
            trial = current + char
            if current and draw.textlength(trial, font=font) > max_width:
                parts.append(current)
                current = char
            else:
                current = trial
        parts.append(current)
    return "\n".join(parts)


def add_sheet(book: Workbook, title: str, name: str) -> None:
    ws = book.create_sheet(title)
    with (OUT / name).open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.reader(fh):
            converted = []
            for value in row:
                try:
                    converted.append(float(value) if value and value[0] in "-+.0123456789" else value)
                except ValueError:
                    converted.append(value)
            ws.append(converted)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="193850")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 36
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            47, max(13, max(len(str(c.value or "")) for c in col[:min(len(col), 80)]) + 2)
        )


def main() -> None:
    assert (OUT / "run_status.txt").read_text().strip() == "0"
    assert (OUT / "publish_status.txt").read_text().strip() == "0"
    algorithms = read("algorithm_strategy_comparison.csv")
    weighted = read("weighted_scheme_comparison.csv")
    assert len(algorithms) == 6 and len(weighted) == 4
    assert all(row["validator"] == "PASS" and int(row["boxes_on_time"]) == 80
               and float(row["weighted_lateness_priority_s"]) == 0
               for row in algorithms + weighted)
    names = [row["scheme"] for row in weighted]
    assert names == ["架次优先", "时间优先", "能耗优先", "综合均衡"]

    algorithm_headers = ["方案", "所用算法", "优化策略", "架次", "总能耗/kWh",
                         "最晚返航/min", "累计飞行交付/h", "按期货箱", "校验"]
    algorithm_rows = [[row["scheme"], row["algorithm"], row["objective_priority"],
                       row["sorties"], f"{float(row['energy_kWh']):.4f}",
                       f"{float(row['Cmax_min']):.4f}",
                       f"{float(row['cumulative_flight_delivery_h']):.4f}",
                       f"{row['boxes_on_time']}/80", row["validator"]]
                      for row in algorithms]
    display_table("算法策略对比.png", "表 1  不同算法与构造策略", 
                  "题目、资源和时间口径一致；逐行重新执行严格 validator。",
                  algorithm_headers, algorithm_rows,
                  [246, 310, 260, 88, 167, 185, 199, 135, 102],
                  "注：历史路径的初始方案与搜索预算不同，本表展示各算法路径的已验证结果，不作为同预算速度竞赛。",
                  highlight_row=5)

    weighted_headers = ["方案", "权重 (架次/时间/能耗)", "架次", "能耗/kWh", "最晚返航/min",
                        "累计飞行交付/h", "加权迟到", "最紧时限裕度/min", "校验"]
    weighted_rows = [[row["scheme"],
                      f"{float(row['sortie_weight']):.2f} / {float(row['time_weight']):.2f} / {float(row['energy_weight']):.2f}",
                      row["sorties"], f"{float(row['energy_kWh']):.4f}",
                      f"{float(row['Cmax_min']):.4f}",
                      f"{float(row['cumulative_flight_delivery_h']):.4f}",
                      f"{float(row['weighted_lateness_priority_s']):.0f}",
                      f"{float(row['minimum_expected_margin_min']):.4f}", row["validator"]]
                     for row in weighted]
    display_table("加权目标方案对比.png", "表 2  加权目标独立求解结果",
                  "四种权重均独立运行、保存排程，并按统一约束复核。",
                  weighted_headers, weighted_rows,
                  [187, 267, 88, 162, 191, 204, 132, 244, 102],
                  "注：80 箱按期是硬约束，故各方案加权迟到为 0；三权重按冻结参考范围归一化后求和。",
                  highlight_row=3)

    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "算法策略对比", "algorithm_strategy_comparison.csv")
    add_sheet(book, "加权目标方案", "weighted_scheme_comparison.csv")
    add_sheet(book, "搜索记录", "actual_strategy_results.csv")
    intro = book.create_sheet("口径说明", 0)
    for row in [
        ["项目", "说明"],
        ["适用问题", "问题二：80 箱、15 个服务区、8 架实体运输无人机、14 组共享电池"],
        ["严格约束", "80/80 箱按期；硬时限、载荷、体积、返航余量、实体机和共享电池时序均通过 validator"],
        ["排程假设", "固定准备可在上一架次起飞后开始；装载须待返航；充电与准备并行，起飞时电池已充满"],
        ["加权目标", "架次、最晚返航、能耗；交付按期是硬约束，故迟到权重不参与可行方案间区分"],
        ["归一化与预算", "冻结参考区间与每个策略 30 秒预算见 scoring_contract.txt；第二轮共享首轮 PASS 候选池"],
        ["算法表限制", "算法路径的历史初始方案与运行预算不同；不得将表1理解为公平速度基准"],
        ["时间列", "最晚返航是全局完成时间；累计飞行交付时长为各架次时长之和，两者不能互换"],
        ["图片来源", "用户参考图片仅提供版式；其中数字没有纳入本工作簿"],
    ]:
        intro.append(row)
    intro.column_dimensions["A"].width = 25
    intro.column_dimensions["B"].width = 105
    for cell in intro[1]:
        cell.fill = PatternFill("solid", fgColor="193850")
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
    path = OUT / "Q2_算法与加权目标方案对比.xlsx"
    book.save(path)
    check = load_workbook(path, read_only=True, data_only=True)
    assert check["算法策略对比"].max_row == 7
    assert check["加权目标方案"].max_row == 5

    a_lines = ["| 方案 | 算法 | 策略 | 架次 | 能耗/kWh | 最晚返航/min | 累计时长/h | 按期 | 校验 |",
               "|---|---|---|---:|---:|---:|---:|---:|---|"]
    a_lines += ["| " + " | ".join(row) + " |" for row in algorithm_rows]
    w_lines = ["| 方案 | 权重 N/T/E | 架次 | 能耗/kWh | 最晚返航/min | 累计时长/h | 加权迟到 | 最紧裕度/min | 校验 |",
               "|---|---:|---:|---:|---:|---:|---:|---:|---|"]
    w_lines += ["| " + " | ".join(row) + " |" for row in weighted_rows]
    best_time = next(row for row in weighted if row["scheme"] == "时间优先")
    best_energy = next(row for row in weighted if row["scheme"] == "能耗优先")
    balanced = next(row for row in weighted if row["scheme"] == "综合均衡")
    text = f"""# 问题二：算法策略与加权目标方案对比

以下两表参考用户图片的**排版结构**，数值均来自本次 MATLAB 实际求解和独立严格 validator；参考图片中的数字未使用。

## 表 1：不同算法与构造策略

{chr(10).join(a_lines)}

表 1 的历史初始方案和预算不同，展示各路径已验证的可行结果；它不是同预算运行速度对照。最后一行采用本轮加权邻域联合搜索的能耗优先排程。`source` 字段列出对应 MAT 文件；到期时间贪心排序方案也已保存为 `due_date_greedy_pass.mat`。

## 表 2：不同目标权重的独立求解

{chr(10).join(w_lines)}

四个分支使用**相同六个严格 PASS 共同种子**、相同多邻域算法和各 30 秒搜索预算，分别保存一份排程。权重为架次／最晚返航／能耗；优化分数使用 `scoring_contract.txt` 所列的冻结参考区间。80 箱按期、医疗与首批硬时限、无人机与电池冲突及安全余量均为硬约束，因此可行解的加权迟到为 0，不另给一个会失去区分力的及时性权重。

时间优先方案最晚返航 **{float(best_time['Cmax_min']):.4f} min**，能耗 **{float(best_time['energy_kWh']):.4f} kWh**；能耗优先方案 **{float(best_energy['energy_kWh']):.4f} kWh**，最晚返航 **{float(best_energy['Cmax_min']):.4f} min**；综合均衡方案 **{float(balanced['Cmax_min']):.4f} min、{float(balanced['energy_kWh']):.4f} kWh**。能耗优先的最紧期望交付裕度仅 **{float(best_energy['minimum_expected_margin_min']):.4f} min**，最小返航安全能量裕度 **{float(best_energy['minimum_energy_reserve_margin_kWh']):.4f} kWh**；虽严格可行，但执行缓冲较小。四方案是当前搜索得到的不同取舍，未证明全局最优。

`Q2_算法与加权目标方案对比.xlsx` 为可编辑表格；`算法策略对比.png`、`加权目标方案对比.png` 可直接作结果图片。四份 `*_best_pass.mat` 是可复核排程，`actual_strategy_results.csv` 和 MATLAB 日志保存搜索记录。
"""
    (OUT / "两阶段方案对比说明.md").write_text(text, encoding="utf-8")
    print(path)


if __name__ == "__main__":
    main()
