"""用统一模型重算方案替换 Q1 提交表，并显式列出每架次作业时间。"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import stat
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


def _copy_style(source_cell, target_cell) -> None:
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def _write_q1_main(sheet, rows: list[dict]) -> None:
    if sheet.max_column < 11 or sheet.cell(1, 8).value != "装卸交接时间（s）":
        sheet.insert_cols(8, 1)
    headers = [
        "架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）", "总体积（m³）",
        "往返时间（s）", "装卸交接时间（s）", "作业时间（s）", "架次能耗（kWh）", "返航SOC（%）",
    ]
    for col, header in enumerate(headers, start=1):
        _copy_style(sheet.cell(1, 7), sheet.cell(1, col))
        sheet.cell(1, col, header)
        sheet.cell(1, col).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    old_max = max(sheet.max_row, len(rows) + 1)
    for row_idx in range(2, old_max + 1):
        for col in range(1, 12):
            sheet.cell(row_idx, col).value = None
    for row_idx, item in enumerate(rows, start=2):
        values = [
            item["batch_id"], item["area_id"], item["uav_type"], item["box_ids_normalized"],
            item["mass_kg"], item["volume_m3"], item["time_s"], item["handling_time_s"],
            item["work_time_s"], item["energy_kwh"], item["soc_percent"],
        ]
        for col, value in enumerate(values, start=1):
            _copy_style(sheet.cell(2, col), sheet.cell(row_idx, col))
            sheet.cell(row_idx, col, value)
            sheet.cell(row_idx, col).alignment = Alignment(
                vertical="top", wrap_text=(col == 4)
            )
        for col in (5, 6, 7, 8, 9, 10):
            sheet.cell(row_idx, col).number_format = "0.00000000"
        sheet.cell(row_idx, 11).number_format = "0.00"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:J{len(rows) + 1}"
    sheet.sheet_view.showGridLines = False
    for col, width in {
        "A": 22, "B": 12, "C": 10, "D": 72, "E": 16, "F": 16,
        "G": 18, "H": 20, "I": 18, "J": 20, "K": 16,
    }.items():
        sheet.column_dimensions[col].width = width


def _write_detail_result(sheet, rows: list[dict]) -> None:
    if sheet.max_column < 12 or sheet.cell(1, 8).value != "装卸交接时间（s）":
        sheet.insert_cols(8, 1)
    headers = [
        "架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）", "总体积（m³）",
        "往返时间（s）", "装卸交接时间（s）", "作业时间（s）", "架次能耗（kWh）", "返航SOC（%）", "箱数",
    ]
    for col, header in enumerate(headers, start=1):
        _copy_style(sheet.cell(1, 7), sheet.cell(1, col))
        sheet.cell(1, col, header)
        sheet.cell(1, col).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    old_max = max(sheet.max_row, len(rows) + 1)
    for row_idx in range(2, old_max + 1):
        for col in range(1, 13):
            sheet.cell(row_idx, col).value = None
    for row_idx, item in enumerate(rows, start=2):
        values = [
            item["batch_id"], item["area_id"], item["uav_type"], item["box_ids_normalized"],
            item["mass_kg"], item["volume_m3"], item["time_s"], item["handling_time_s"],
            item["work_time_s"], item["energy_kwh"], item["soc_percent"], item["box_count"],
        ]
        for col, value in enumerate(values, start=1):
            _copy_style(sheet.cell(2, col), sheet.cell(row_idx, col))
            sheet.cell(row_idx, col, value)
            sheet.cell(row_idx, col).alignment = Alignment(
                vertical="top", wrap_text=(col == 4)
            )
        for col in (5, 6, 7, 8, 9, 10):
            sheet.cell(row_idx, col).number_format = "0.00000000"
        sheet.cell(row_idx, 11).number_format = "0.00"
        sheet.cell(row_idx, 12).number_format = "0"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{len(rows) + 1}"
    sheet.sheet_view.showGridLines = False
    for col, width in {
        "A": 22, "B": 12, "C": 10, "D": 72, "E": 16, "F": 16,
        "G": 18, "H": 20, "I": 18, "J": 20, "K": 16, "L": 10,
    }.items():
        sheet.column_dimensions[col].width = width


def replace_q1_submission(
    source_workbook: Path,
    result_json: Path,
    output_workbook: Path,
) -> None:
    result = json.loads(result_json.read_text(encoding="utf-8"))
    rows = result["per_batch"]
    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    temp_path = output_workbook.with_name(output_workbook.stem + ".tmp.xlsx")
    if temp_path.exists():
        temp_path.unlink()
    shutil.copy2(source_workbook, temp_path)
    os.chmod(temp_path, stat.S_IREAD | stat.S_IWRITE)
    workbook = load_workbook(temp_path, data_only=False)
    _write_q1_main(workbook["Q1_单点组批"], rows)
    _write_detail_result(workbook["Q1_统一模型结果"], rows)

    notes = workbook["填写说明"]
    row_idx = notes.max_row + 1
    notes.cell(row_idx, 1, "Q1 主表更新")
    notes.cell(
        row_idx,
        2,
        "已替换为统一模型重算的 18 架次方案；保留“往返时间（s）”，并将“作业时间（s）”定义为"
        "往返飞行 + 工位固定准备 + 每箱装载 + 接收点基础交接 + 每箱增加交接。"
        "时间参数来自运输无人机数据.xlsx。",
    )
    notes.cell(row_idx, 1).font = copy(notes.cell(5, 1).font)
    notes.cell(row_idx, 1).fill = copy(notes.cell(5, 1).fill)
    notes.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    notes.column_dimensions["A"].width = max(notes.column_dimensions["A"].width or 0, 28)
    notes.column_dimensions["B"].width = max(notes.column_dimensions["B"].width or 0, 110)
    workbook.active = workbook.sheetnames.index("Q1_单点组批")
    workbook.save(temp_path)
    if output_workbook.exists():
        os.chmod(output_workbook, stat.S_IREAD | stat.S_IWRITE)
    os.replace(temp_path, output_workbook)


def main() -> int:
    parser = argparse.ArgumentParser(description="替换结果提交模板中的 Q1 并增加作业时间列")
    parser.add_argument(
        "--source-workbook",
        type=Path,
        default=Path("results/问题1_对照方案_统一模型.xlsx"),
    )
    parser.add_argument(
        "--result-json",
        type=Path,
        default=Path("results/q1_对照方案_统一模型.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("results/结果提交模板_Q1_已填写.xlsx"),
    )
    args = parser.parse_args()
    replace_q1_submission(args.source_workbook, args.result_json, args.output)
    print(f"created {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
