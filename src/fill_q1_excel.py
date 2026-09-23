"""Fill the supplied result template with the verified Q1 outputs.

The input template and raw attachments are read-only.  This module creates a
new workbook under ``results/`` and keeps the untouched Q2--Q4 template
worksheets so later team members can fill them independently.
"""

from __future__ import annotations

import argparse
import csv
import json
from copy import copy
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.data_loader import load_inputs


Q1_TEMPLATE_HEADERS = [
    "架次编号",
    "服务区编号",
    "机型编号",
    "货箱编号列表",
    "总质量（kg）",
    "总体积（m³）",
    "往返时间（s）",
    "作业时间（s）",
    "架次能耗（kWh）",
    "返航SOC（%）",
]


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _float(value: Any) -> float:
    return float(value)


def _yes_no(value: Any) -> str:
    return "是" if bool(value) else "否"


def _copy_row_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _set_widths(ws, headers: list[str], rows: list[list[Any]]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        values = [header]
        values.extend("" if row[col_idx - 1] is None else str(row[col_idx - 1]) for row in rows)
        width = min(max(len(value) for value in values) + 2, 52)
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else f"A{col_idx}"].width = width


def _format_table(ws, headers: list[str], rows: list[list[Any]], table_name: str) -> None:
    """Write a filterable plain range instead of an OOXML structured table.

    The supplied template is intended for broad Excel/WPS compatibility.  A
    plain range with an autofilter avoids structured-table repair prompts in
    older Excel-compatible clients while retaining the same comparison data.
    ``table_name`` remains an argument so callers document the logical table
    identity without creating a fragile table relationship.
    """
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.alignment = Alignment(vertical="top", wrap_text=col_idx in {4, 5})
            header = headers[col_idx - 1]
            if header.endswith("_%") or "（%）" in header:
                cell.number_format = "0.00"
            elif any(unit in header for unit in ("_kg", "_m3", "_m", "_s", "_kWh", "（kg）", "（m³）", "（s）", "（kWh）")):
                cell.number_format = "0.000"
    _set_widths(ws, headers, rows)
    ws.auto_filter.ref = f"A1:{chr(64 + len(headers)) if len(headers) <= 26 else 'A' + str(len(headers))}{len(rows) + 1}"


def _new_table_sheet(workbook, title: str, headers: list[str], rows: list[list[Any]], table_name: str):
    if title in workbook.sheetnames:
        del workbook[title]
    ws = workbook.create_sheet(title)
    _format_table(ws, headers, rows, table_name)
    return ws


def _clear_q1_template_sheet(ws) -> None:
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), max_col=len(Q1_TEMPLATE_HEADERS)):
        for cell in row:
            cell.value = None


def _fill_q1_template_sheet(ws, batch_rows: list[dict[str, str]]) -> None:
    if ws.max_column < len(Q1_TEMPLATE_HEADERS) or ws.cell(1, 8).value != "作业时间（s）":
        ws.insert_cols(8, 1)
    headers = [ws.cell(1, col).value for col in range(1, len(Q1_TEMPLATE_HEADERS) + 1)]
    if headers != Q1_TEMPLATE_HEADERS:
        raise ValueError(f"Q1 模板表头不匹配: {headers!r}")
    _clear_q1_template_sheet(ws)
    style_source_row = 2
    for row_idx, record in enumerate(batch_rows, start=2):
        for col_idx in range(1, len(Q1_TEMPLATE_HEADERS) + 1):
            _copy_row_style(ws.cell(style_source_row, col_idx), ws.cell(row_idx, col_idx))
        values = [
            record["架次编号"],
            record["服务区编号"],
            record["机型编号"],
            record["货箱编号列表"],
            _float(record["总质量_kg"]),
            _float(record["总体积_m3"]),
            _float(record["往返时间_s"]),
            _float(record["作业时间_s"]),
            _float(record["架次能耗_kWh"]),
            _float(record["返航SOC_%"]),
        ]
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row_idx, col_idx, value)
        ws.cell(row_idx, 4).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row_idx, 5).number_format = "0.000"
        ws.cell(row_idx, 6).number_format = "0.000"
        ws.cell(row_idx, 7).number_format = "0.000"
        ws.cell(row_idx, 8).number_format = "0.000"
        ws.cell(row_idx, 9).number_format = "0.000"
        ws.cell(row_idx, 10).number_format = "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{len(batch_rows) + 1}"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["D"].width = 56


def _write_notes_sheet(workbook, rows: list[tuple[str, str]]) -> None:
    if "填写说明" in workbook.sheetnames:
        del workbook["填写说明"]
    ws = workbook.create_sheet("填写说明", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    ws.merge_cells("A1:B1")
    ws["A1"] = "D题问题一结果提交与详细对比"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].alignment = Alignment(horizontal="center")
    for row_idx, (key, value) in enumerate(rows, start=3):
        ws.cell(row_idx, 1, key).font = Font(bold=True)
        ws.cell(row_idx, 1).fill = PatternFill("solid", fgColor="D9EAF7")
        ws.cell(row_idx, 2, value)
        ws.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"


def build_workbook(template_path: Path, data_dir: Path, results_dir: Path, output_path: Path) -> None:
    batch_rows = _read_csv(results_dir / "q1_组批方案.csv")
    method_rows = _read_csv(results_dir / "q1_方法比较.csv")
    safe_rows = _read_csv(results_dir / "q1_安全载荷.csv")
    route_rows = _read_csv(results_dir / "q1_路线几何.csv")
    sensitivity = json.loads((results_dir / "q1_余量敏感性.json").read_text(encoding="utf-8"))
    validation = json.loads((results_dir / "q1_约束验证.json").read_text(encoding="utf-8"))
    run_status = json.loads((results_dir / "运行状态.json").read_text(encoding="utf-8"))

    inputs = load_inputs(data_dir)
    uav_types = inputs.uav_types
    safe_by_key = {(row["机型编号"], row["服务区编号"]): row for row in safe_rows}
    route_by_area = {row["服务区编号"]: row for row in route_rows}

    detail_headers = [
        "架次编号", "服务区编号", "机型编号", "货箱数", "货箱编号列表",
        "总质量_kg", "安全载荷_kg", "质量余量_kg", "质量利用率_%",
        "总体积_m3", "机型体积上限_m3", "体积余量_m3", "体积利用率_%",
        "单程距离_m", "往返距离_m", "安全载荷对应航程_m", "航程余量_m",
        "架次能耗_kWh", "允许任务能量_kWh", "能耗占允许_%", "往返时间_s", "装卸交接时间_s", "作业时间_s", "返航SOC_%",
    ]
    detail_data: list[list[Any]] = []
    for record in batch_rows:
        area = record["服务区编号"]
        type_id = record["机型编号"]
        safe = safe_by_key[(type_id, area)]
        route = route_by_area[area]
        uav = uav_types[type_id]
        mass = _float(record["总质量_kg"])
        volume = _float(record["总体积_m3"])
        safe_payload = _float(safe["最大安全载荷_kg"])
        volume_limit = float(uav.volume_capacity_m3)
        roundtrip_distance = _float(safe["往返距离_m"])
        range_at_safe = _float(safe["安全载荷对应航程_m"])
        energy = _float(record["架次能耗_kWh"])
        allowed_energy = _float(safe["允许任务能量_kWh"])
        detail_data.append([
            record["架次编号"], area, type_id,
            len(record["货箱编号列表"].split(";")), record["货箱编号列表"],
            mass, safe_payload, safe_payload - mass, 100.0 * mass / safe_payload,
            volume, volume_limit, volume_limit - volume, 100.0 * volume / volume_limit,
            _float(route["单程距离_m"]), roundtrip_distance, range_at_safe,
            range_at_safe - roundtrip_distance, energy, allowed_energy,
            100.0 * energy / allowed_energy, _float(record["往返时间_s"]),
            _float(record["装卸交接时间_s"]), _float(record["作业时间_s"]),
            _float(record["返航SOC_%"]),
        ])

    method_key = {
        "FFD baseline": "ffd_baseline",
        "local search": "local_search",
        "MILP minimum-sortie": "milp",
        "MILP + local search": "milp_local_search",
    }
    base = method_rows[0]
    base_sorties = _float(base["架次数"])
    base_energy = _float(base["总运输能耗_kWh"])
    base_time = _float(base["累计作业时间_s"])
    comparison_headers = [
        "方法", "架次数", "架次节省_架", "架次节省_%", "总运输能耗_kWh",
        "能耗节省_kWh", "能耗节省_%", "累计作业时间_s", "时间节省_s", "时间节省_%",
        "搜索是否完整", "约束验证", "说明",
    ]
    comparison_data: list[list[Any]] = []
    for row in method_rows:
        sorties = _float(row["架次数"])
        energy = _float(row["总运输能耗_kWh"])
        time_s = _float(row["累计作业时间_s"])
        check = validation[method_key[row["方法"]]]
        if row["方法"] == "FFD baseline":
            note = "确定性首个可行架次基线，不宣称最优"
        elif row["方法"] == "local search":
            note = "同服务区 relocate/swap，字典序目标为架次、能耗、时间"
        elif row["方法"] == "MILP minimum-sortie":
            note = "质量/体积/安全载荷容量 MILP；30 s 内搜索完整"
        else:
            note = "从 MILP 解出发做局部搜索；改善能耗但不宣称全局最优"
        comparison_data.append([
            row["方法"], int(sorties), int(base_sorties - sorties),
            100.0 * (base_sorties - sorties) / base_sorties,
            energy, base_energy - energy, 100.0 * (base_energy - energy) / base_energy,
            time_s, base_time - time_s, 100.0 * (base_time - time_s) / base_time,
            _yes_no(row["搜索是否完整"] == "True"), _yes_no(check["passed"]), note,
        ])

    safe_headers = [
        "机型编号", "服务区编号", "返航余量ρ", "最大安全载荷_kg", "额定最大载荷_kg",
        "安全载荷占额定_%", "往返能耗_kWh", "允许任务能量_kWh", "能量余量_kWh",
        "往返距离_m", "安全载荷对应航程_m", "航程余量_m", "零载荷是否可行", "搜索方法",
    ]
    safe_data: list[list[Any]] = []
    for row in safe_rows:
        uav = uav_types[row["机型编号"]]
        allowed = _float(row["允许任务能量_kWh"])
        usable = float(uav.usable_energy_kwh)
        safe_payload = _float(row["最大安全载荷_kg"])
        safe_data.append([
            row["机型编号"], row["服务区编号"], 1.0 - allowed / usable,
            safe_payload, float(uav.max_payload_kg), 100.0 * safe_payload / float(uav.max_payload_kg),
            _float(row["往返能耗_kWh"]), allowed, allowed - _float(row["往返能耗_kWh"]),
            _float(row["往返距离_m"]), _float(row["安全载荷对应航程_m"]),
            _float(row["安全载荷对应航程_m"]) - _float(row["往返距离_m"]),
            _yes_no(row["零载荷是否可行"] == "True"), row["搜索方法"],
        ])

    switch_by_to = {float(item["to_rho"]): ", ".join(item["changed_metrics"]) for item in sensitivity["switch_points"]}
    sensitivity_headers = [
        "返航余量ρ", "FFD架次", "局部搜索架次", "架次节省_架", "FFD能耗_kWh",
        "局部搜索能耗_kWh", "局部搜索相对FFD节能_%", "FFD时间_s", "局部搜索时间_s",
        "A型最小安全载荷_kg", "B型最小安全载荷_kg", "C型最小安全载荷_kg",
        "A型受限服务区数", "B型受限服务区数", "C型受限服务区数", "相对上一ρ变化",
    ]
    sensitivity_data: list[list[Any]] = []
    previous = None
    for record in sensitivity["records"]:
        rho = float(record["rho"])
        payloads = record["safe_payloads_kg"]
        ffd = record["ffd"]
        local = record["local_search"]
        values = [
            rho, int(ffd["sorties"]), int(local["sorties"]), int(ffd["sorties"] - local["sorties"]),
            float(ffd["energy_kwh"]), float(local["energy_kwh"]),
            100.0 * (float(ffd["energy_kwh"]) - float(local["energy_kwh"])) / float(ffd["energy_kwh"]),
            float(ffd["time_s"]), float(local["time_s"]),
        ]
        minimum_payloads = []
        restricted_counts = []
        for type_id in ("A", "B", "C"):
            type_values = [float(value) for key, value in payloads.items() if key.startswith(f"{type_id}/")]
            max_payload = float(uav_types[type_id].max_payload_kg)
            minimum_payloads.append(min(type_values))
            restricted_counts.append(sum(value < max_payload - 1e-8 for value in type_values))
        values.extend(minimum_payloads + restricted_counts)
        if previous is None:
            change = "首个扫描点"
        else:
            changed = switch_by_to.get(rho, "")
            change = changed or "无架次/安全载荷指标变化"
        sensitivity_data.append(values + [change])
        previous = record

    constraint_headers = [
        "方法", "架次数", "货箱数", "搜索是否完整", "总体通过", "箱体唯一覆盖", "不跨服务区",
        "质量容量", "体积容量", "能耗余量", "航程约束", "总量一致", "有限数值", "备注",
    ]
    constraint_data: list[list[Any]] = []
    for label, key in (("FFD baseline", "ffd_baseline"), ("local search", "local_search"), ("MILP minimum-sortie", "milp"), ("MILP + local search", "milp_local_search")):
        item = validation[key]
        checks = item["checks"]
        constraint_data.append([
            label, item["n_batches"], item["n_boxes"], _yes_no(item["complete_search"]), _yes_no(item["passed"]),
            _yes_no(checks["all_boxes_present_once"]), _yes_no(checks["no_cross_area_batch"]),
            _yes_no(checks["mass_capacity"]), _yes_no(checks["volume_capacity"]),
            _yes_no(checks["energy_reserve"]), _yes_no(checks["range_constraint"]),
            _yes_no(checks["solution_totals_match"]), _yes_no(checks["finite_units"]),
            "; ".join(item.get("notes", [])) or "无",
        ])

    workbook = load_workbook(template_path, data_only=False)
    if "Q1_单点组批" not in workbook.sheetnames:
        raise ValueError("结果模板缺少 Q1_单点组批 工作表")
    _fill_q1_template_sheet(workbook["Q1_单点组批"], batch_rows)

    notes = [
        ("文件用途", "基于题目提供的结果提交模板，仅填写问题一；Q2-Q4 保留原模板表头和空白区域。"),
        ("输入 ZIP SHA-256", "BBB4C17705B70B1BA365523AB189810305DB86D6457AED5F4F5587274E4D07B4"),
        ("计算范围", "固定服务区内 O01→S_i→O01 单点往返；不可拆货箱；允许同一服务区多架次；不跨服务区组批。"),
        ("选定方案", f"{run_status['selected_method']}；{run_status['selected_metrics']['sorties']} 架次；{run_status['selected_metrics']['energy_kwh']:.8f} kWh；{run_status['selected_metrics']['time_s']:.8f} s。"),
        ("能耗模型状态", "E_hor 与 E_up 子公式未在题面展开；当前使用 src.problem_energy_provider.py 登记的 reference-model-A-v1，表内数值为条件性结果。"),
        ("比较口径", "架次/能耗/时间节省均相对 FFD baseline 计算，节省为正；MILP 仅对架次最小化，MILP+local search 不宣称全局最优。"),
        ("Skill 版本", "math-modeling 1.3.0；源仓库 commit 2b0322483699e87a7afecdd3fe4c4cb4858681be。"),
        ("原始附件", "data/raw/D题 下的快照已设为只读；本工作簿另存于 results，不修改原始模板。"),
    ]
    _write_notes_sheet(workbook, notes)
    _new_table_sheet(workbook, "Q1_架次明细", detail_headers, detail_data, "Q1BatchDetail")
    _new_table_sheet(workbook, "Q1_方法对比", comparison_headers, comparison_data, "Q1MethodCompare")
    _new_table_sheet(workbook, "Q1_安全载荷", safe_headers, safe_data, "Q1SafePayload")
    _new_table_sheet(workbook, "Q1_敏感性", sensitivity_headers, sensitivity_data, "Q1Sensitivity")
    _new_table_sheet(workbook, "Q1_约束核验", constraint_headers, constraint_data, "Q1Validation")
    _new_table_sheet(workbook, "Q1_路线几何", list(route_rows[0].keys()), [[row[key] for key in route_rows[0].keys()] for row in route_rows], "Q1RouteGeometry")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="填写问题一结果模板并生成详细对比工作表")
    parser.add_argument("--template", type=Path, default=Path("data/raw/D题/结果提交模板.xlsx"))
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--results-dir", type=Path, default=Path("results"))
    parser.add_argument("--output", type=Path, default=Path("results/结果提交模板_Q1_已填写.xlsx"))
    args = parser.parse_args()
    build_workbook(args.template, args.data_dir, args.results_dir, args.output)
    print(f"created {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
