"""按统一参考模型重算外部提供的问题一组批方案。

本脚本把外部工作簿当作只读的“方案输入”，不猜测它原先采用的公式；
逐架次使用项目登记的 ``reference-model-A-v1`` 重算时间、能耗、返航
SOC 和约束余量，并生成机器可读结果与兼容性较好的普通 Excel 工作簿。

运行示例::

    python -m src.recompute_q1_solution \
      --solution data/raw/D题/对照方案_问题1_结果提交.xlsx \
      --data-dir data --output-dir results

注意：题面只展开了 E=E_hor+E_up，没有给出两个子项的唯一展开式；
因此本脚本的结果是当前项目参考模型下的条件性结果，不能冒充官方唯一答案。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import shutil
import stat
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from .data_loader import load_inputs
from .energy_model import (
    equivalent_range_m,
    handling_time_s,
    load_energy_provider,
    round_trip_energy_kwh,
    round_trip_time_s,
    sortie_work_time_s,
)
from .payload_solver import max_safe_payload
from .terrain import DEM, build_routes


SOLUTION_SHEET = "Q1_单点组批"
BOX_SPLIT_RE = re.compile(r"[、，,;；\s]+")
TOLERANCE = 1e-7


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _json_number(value: float) -> float:
    """Return a finite float suitable for JSON and validation output."""

    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"非有限数值: {value!r}")
    return result


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[\s\u3000_\-（）()［］\[\]：:，,。./]", "", text).lower()


def _text(value: Any, *, field: str, row: int) -> str:
    if value is None:
        raise ValueError(f"外部方案第 {row} 行字段 {field} 为空")
    text = str(value).strip()
    if not text:
        raise ValueError(f"外部方案第 {row} 行字段 {field} 为空")
    return text


def _number(value: Any, *, field: str, row: int) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"外部方案第 {row} 行字段 {field} 为空")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"外部方案第 {row} 行字段 {field} 不是数值: {value!r}"
        ) from exc
    if not math.isfinite(result):
        raise ValueError(f"外部方案第 {row} 行字段 {field} 不是有限数")
    return result


def _find_column(headers: Iterable[Any], aliases: Iterable[str], *, table: str) -> int:
    headers = tuple(headers)
    normalized = [_norm(item) for item in headers]
    alias_norm = tuple(_norm(item) for item in aliases)
    exact = [idx for idx, value in enumerate(normalized) if value in alias_norm]
    if len(exact) == 1:
        return exact[0]
    contains = [
        idx
        for idx, value in enumerate(normalized)
        if value and any(alias and (alias in value or value in alias) for alias in alias_norm)
    ]
    if len(contains) == 1:
        return contains[0]
    raise ValueError(
        f"{table} 无法唯一定位字段 {list(aliases)!r}；实际表头为 {list(headers)!r}"
    )


def _parse_box_ids(value: Any, *, row: int) -> tuple[str, ...]:
    text = _text(value, field="货箱编号列表", row=row)
    box_ids = tuple(item.strip() for item in BOX_SPLIT_RE.split(text) if item.strip())
    if not box_ids:
        raise ValueError(f"外部方案第 {row} 行没有可解析的货箱编号")
    return box_ids


def _load_solution_rows(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取外部 XLSX 需要 openpyxl") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    if SOLUTION_SHEET in workbook.sheetnames:
        sheet = workbook[SOLUTION_SHEET]
    else:
        q1_sheets = [name for name in workbook.sheetnames if _norm(name).startswith("q1")]
        if len(q1_sheets) != 1:
            raise ValueError(
                f"外部方案缺少明确的 {SOLUTION_SHEET} 工作表，候选为 {q1_sheets!r}"
            )
        sheet = workbook[q1_sheets[0]]

    rows = list(sheet.iter_rows(values_only=True))
    sheetnames = list(workbook.sheetnames)
    if not rows:
        raise ValueError(f"外部方案工作表为空: {path} / {sheet.title}")
    headers = tuple(rows[0])
    table = f"{path.name}/{sheet.title}"
    columns = {
        "batch_id": _find_column(headers, ("架次编号", "批次编号", "batch_id"), table=table),
        "area_id": _find_column(headers, ("服务区编号", "服务区", "area_id"), table=table),
        "uav_type": _find_column(headers, ("机型编号", "机型", "无人机型号", "type_id"), table=table),
        "box_ids": _find_column(headers, ("货箱编号列表", "货箱编号", "箱体编号", "box_ids"), table=table),
        "mass": _find_column(headers, ("总质量（kg）", "总质量(kg)", "总质量", "mass_kg"), table=table),
        "volume": _find_column(headers, ("总体积（m³）", "总体积(m³)", "总体积", "volume_m3"), table=table),
        "time": _find_column(headers, ("往返时间（s）", "往返时间(s)", "往返时间", "time_s"), table=table),
        "energy": _find_column(headers, ("架次能耗（kWh）", "架次能耗(kWh)", "架次能耗", "energy_kwh"), table=table),
        "soc": _find_column(headers, ("返航SOC（%）", "返航SOC(%)", "返航SOC", "soc"), table=table),
    }

    parsed: list[dict[str, Any]] = []
    for row_number, raw in enumerate(rows[1:], start=2):
        if not any(item is not None and str(item).strip() for item in raw):
            continue
        parsed.append(
            {
                "source_row": row_number,
                "batch_id": _text(raw[columns["batch_id"]], field="架次编号", row=row_number),
                "area_id": _text(raw[columns["area_id"]], field="服务区编号", row=row_number).upper(),
                "uav_type": _text(raw[columns["uav_type"]], field="机型编号", row=row_number).upper(),
                "box_ids": _parse_box_ids(raw[columns["box_ids"]], row=row_number),
                "source_mass_kg": _number(raw[columns["mass"]], field="总质量", row=row_number),
                "source_volume_m3": _number(raw[columns["volume"]], field="总体积", row=row_number),
                "source_time_s": _number(raw[columns["time"]], field="往返时间", row=row_number),
                "source_energy_kwh": _number(raw[columns["energy"]], field="架次能耗", row=row_number),
                "source_soc_percent": _number(raw[columns["soc"]], field="返航SOC", row=row_number),
            }
        )
    workbook.close()
    if not parsed:
        raise ValueError(f"外部方案 {path} 没有可读取的 Q1 架次行")
    return parsed, sheetnames


def _relative_or_absolute(path: Path, root: Path) -> str:
    try:
        return str(path.resolve().relative_to(root.resolve()))
    except ValueError:
        return str(path.resolve())


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _style_sheet(sheet: Any, *, freeze: str = "A2", widths: dict[int, float] | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 32
    if widths:
        for column_index, width in widths.items():
            sheet.column_dimensions[chr(64 + column_index) if column_index <= 26 else "A"].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def _set_column_widths(sheet: Any, widths: list[float]) -> None:
    from openpyxl.utils import get_column_letter

    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width


def _write_workbook(
    output_path: Path,
    source_path: Path,
    detail_rows: list[dict[str, Any]],
    validation: dict[str, Any],
    metrics: dict[str, Any],
    model: dict[str, Any],
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)
    # The source snapshot is deliberately read-only; the generated copy must be writable.
    output_path.chmod(stat.S_IREAD | stat.S_IWRITE)
    workbook = load_workbook(output_path, data_only=False)

    for name in ("填写说明", "Q1_统一模型结果", "Q1_原表对比", "Q1_统一模型核验"):
        if name in workbook.sheetnames:
            del workbook[name]

    notes = workbook.create_sheet("填写说明", 0)
    notes.append(["项目", "统一模型重算外部 Q1 方案"])
    notes.append(["输入快照", _relative_or_absolute(source_path, output_path.parent.parent)])
    notes.append(["模型", model["provider_source"]])
    notes.append(["适用范围", model["scope"]])
    notes.append(["重算架次", metrics["recomputed"]["sorties"]])
    notes.append(["重算总能耗_kWh", metrics["recomputed"]["energy_kwh"]])
    notes.append(["重算往返飞行时间_s", metrics["recomputed"]["time_s"]])
    notes.append(["装卸交接总时间_s", metrics["recomputed"]["handling_time_s"]])
    notes.append(["重算总作业时间_s", metrics["recomputed"]["work_time_s"]])
    notes.append(["相对原表能耗差_kWh", metrics["delta"]["energy_kwh"]])
    notes.append(["相对原表时间差_s", metrics["delta"]["time_s"]])
    notes.append(["约束核验", "通过" if validation["passed"] else "未通过"])
    notes.append(["说明", model["caveat"]])
    _style_sheet(notes, freeze="A2")
    _set_column_widths(notes, [28, 110])
    for row in notes.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    result = workbook.create_sheet("Q1_统一模型结果")
    result_headers = [
        "架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）", "总体积（m³）",
        "往返时间（s）", "装卸交接时间（s）", "作业时间（s）", "架次能耗（kWh）", "返航SOC（%）", "箱数",
    ]
    result.append(result_headers)
    for row in detail_rows:
        result.append([
            row["batch_id"], row["area_id"], row["uav_type"], row["box_ids_normalized"],
            row["mass_kg"], row["volume_m3"], row["time_s"], row["handling_time_s"],
            row["work_time_s"], row["energy_kwh"], row["soc_percent"], row["box_count"],
        ])
    _style_sheet(result, freeze="A2")
    _set_column_widths(result, [22, 12, 10, 72, 16, 16, 18, 20, 18, 20, 16, 10])
    for row in result.iter_rows(min_row=2, min_col=5, max_col=10):
        for cell in row:
            cell.number_format = "0.00000000"

    comparison = workbook.create_sheet("Q1_原表对比")
    comparison_headers = [
        "架次编号", "服务区编号", "机型编号", "货箱编号列表", "原表总质量_kg", "统一模型总质量_kg",
        "质量差_kg", "原表总体积_m3", "统一模型总体积_m3", "体积差_m3", "原表往返时间_s",
        "统一模型往返时间_s", "时间差_s", "统一模型作业时间_s", "原表架次能耗_kWh", "统一模型架次能耗_kWh",
        "能耗差_kWh", "原表返航SOC_%", "统一模型返航SOC_%", "SOC差_百分点", "安全载荷_kg",
        "质量余量_kg", "体积余量_m3", "往返距离_m", "等效航程_m", "航程余量_m", "允许任务能量_kWh",
        "能量余量_kWh", "源数据质量体积一致", "约束核验通过", "备注",
    ]
    comparison.append(comparison_headers)
    for row in detail_rows:
        comparison.append([
            row["batch_id"], row["area_id"], row["uav_type"], row["box_ids_normalized"],
            row["source_mass_kg"], row["mass_kg"], row["mass_diff_kg"],
            row["source_volume_m3"], row["volume_m3"], row["volume_diff_m3"],
            row["source_time_s"], row["time_s"], row["time_diff_s"], row["work_time_s"],
            row["source_energy_kwh"], row["energy_kwh"], row["energy_diff_kwh"],
            row["source_soc_percent"], row["soc_percent"], row["soc_diff_percentage_points"],
            row["safe_payload_kg"], row["mass_slack_kg"], row["volume_slack_m3"],
            row["round_trip_distance_m"], row["equivalent_range_m"], row["range_slack_m"],
            row["allowed_energy_kwh"], row["energy_slack_kwh"], row["source_mass_volume_match"],
            row["constraint_passed"], row["备注"],
        ])
    _style_sheet(comparison, freeze="A2")
    _set_column_widths(comparison, [22, 12, 10, 72] + [18] * (len(comparison_headers) - 4))
    for row in comparison.iter_rows(min_row=2, min_col=5, max_col=len(comparison_headers) - 2):
        for cell in row:
            cell.number_format = "0.00000000"

    check_sheet = workbook.create_sheet("Q1_统一模型核验")
    check_sheet.append(["核验项目", "结果", "说明"])
    checks = validation["checks"]
    check_rows = [
        ("总体通过", "通过" if validation["passed"] else "未通过", "; ".join(validation["errors"]) or "无错误"),
        ("架次数", validation["n_batches"], "外部方案非空架次数量"),
        ("货箱数", validation["n_boxes"], "题目输入货箱数量"),
        ("箱体唯一覆盖", "通过" if checks["all_boxes_present_once"] else "未通过", "80 个货箱各出现一次"),
        ("不跨服务区", "通过" if checks["no_cross_area_batch"] else "未通过", "每架次只包含本服务区箱体"),
        ("质量容量", "通过" if checks["mass_capacity"] else "未通过", "质量不超过安全载荷"),
        ("体积容量", "通过" if checks["volume_capacity"] else "未通过", "体积不超过机型舱容"),
        ("能耗返航余量", "通过" if checks["energy_reserve"] else "未通过", "能耗不超过允许任务能量"),
        ("等效航程", "通过" if checks["range_constraint"] else "未通过", "往返距离不超过当前载荷等效航程"),
        ("源数据质量/体积一致", "通过" if checks["source_mass_volume_match"] else "未通过", "原表聚合值与输入箱体求和一致"),
        ("有限数值", "通过" if checks["finite_units"] else "未通过", "输出数值均为有限数"),
        ("原表总能耗_kWh", metrics["source_reported"]["energy_kwh"], "由外部方案逐架次能耗求和"),
        ("统一模型总能耗_kWh", metrics["recomputed"]["energy_kwh"], "reference-model-A-v1 重算"),
        ("原表总往返时间_s", metrics["source_reported"]["time_s"], "由外部方案逐架次往返飞行时间求和"),
        ("统一模型总往返飞行时间_s", metrics["recomputed"]["time_s"], "按项目路线几何和机型速度重算"),
        ("统一模型装卸交接总时间_s", metrics["recomputed"]["handling_time_s"], "按运输无人机附件的准备、装载和接收点交接参数求和"),
        ("统一模型总作业时间_s", metrics["recomputed"]["work_time_s"], "往返飞行时间 + 装卸交接时间"),
    ]
    for item in check_rows:
        check_sheet.append(item)
    _style_sheet(check_sheet, freeze="A2")
    _set_column_widths(check_sheet, [32, 24, 110])
    for row in check_sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "0.00000000"
    # Make the summary visually distinct without introducing Excel Table XML.
    for cell in check_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    workbook.active = workbook.sheetnames.index("Q1_统一模型结果")
    workbook.save(output_path)


def recompute(
    solution_path: Path,
    data_dir: Path,
    output_dir: Path,
    excel_output: Path,
    energy_provider_module: str,
) -> dict[str, Any]:
    solution_path = solution_path.resolve()
    data_dir = data_dir.resolve()
    output_dir = output_dir.resolve()
    excel_output = excel_output.resolve()
    rows, sheetnames = _load_solution_rows(solution_path)

    inputs = load_inputs(data_dir)
    provider = load_energy_provider(energy_provider_module, inputs.uav_types)
    box_by_id = {box.box_id: box for box in inputs.boxes}

    with DEM.open(inputs.dem_path) as dem:
        routes = build_routes(inputs.nodes, dem)

    safe_payloads: dict[tuple[str, str], float] = {}
    for type_id, uav in inputs.uav_types.items():
        for area_id, route in routes.items():
            safe_payloads[(type_id, area_id)] = max_safe_payload(
                provider, uav, route
            ).safe_payload_kg

    errors: list[str] = []
    seen: list[str] = []
    detail_rows: list[dict[str, Any]] = []
    checks = {
        "all_boxes_present_once": True,
        "no_cross_area_batch": True,
        "mass_capacity": True,
        "volume_capacity": True,
        "energy_reserve": True,
        "range_constraint": True,
        "source_mass_volume_match": True,
        "finite_units": True,
    }

    for item in rows:
        batch_id = item["batch_id"]
        area_id = item["area_id"]
        type_id = item["uav_type"]
        box_ids = item["box_ids"]
        seen.extend(box_ids)
        unknown = [box_id for box_id in box_ids if box_id not in box_by_id]
        if unknown:
            checks["all_boxes_present_once"] = False
            errors.append(f"{batch_id} 含未知货箱: {unknown}")
        if area_id not in routes:
            checks["no_cross_area_batch"] = False
            errors.append(f"{batch_id} 的服务区不存在: {area_id}")
            continue
        if type_id not in inputs.uav_types:
            checks["finite_units"] = False
            errors.append(f"{batch_id} 的机型不存在: {type_id}")
            continue
        boxes = [box_by_id[box_id] for box_id in box_ids if box_id in box_by_id]
        box_areas = {box.service_area_id for box in boxes}
        if box_areas != {area_id} or len(boxes) != len(box_ids):
            checks["no_cross_area_batch"] = False
            errors.append(f"{batch_id} 跨服务区或含未知箱体")

        uav = inputs.uav_types[type_id]
        route = routes[area_id]
        mass = sum(box.mass_kg for box in boxes)
        volume = sum(box.volume_m3 for box in boxes)
        safe_payload = safe_payloads[(type_id, area_id)]
        allowed = (1.0 - uav.reserve_ratio) * uav.usable_energy_kwh
        distance = route.outbound.distance_m + route.inbound.distance_m
        range_m = equivalent_range_m(uav, mass)
        energy = round_trip_energy_kwh(provider, uav, route, mass)
        time_s = round_trip_time_s(uav, route)
        handling_s = handling_time_s(uav, len(box_ids))
        work_time_s = sortie_work_time_s(uav, route, len(box_ids))
        soc_percent = 100.0 * (1.0 - energy / uav.usable_energy_kwh)

        mass_match = abs(item["source_mass_kg"] - mass) <= 1e-7
        volume_match = abs(item["source_volume_m3"] - volume) <= 1e-9
        if not (mass_match and volume_match):
            checks["source_mass_volume_match"] = False
            errors.append(f"{batch_id} 原表质量/体积与输入箱体求和不一致")
        if mass > safe_payload + TOLERANCE:
            checks["mass_capacity"] = False
            errors.append(f"{batch_id} 超过安全载荷")
        if volume > uav.volume_capacity_m3 + TOLERANCE:
            checks["volume_capacity"] = False
            errors.append(f"{batch_id} 超过装载体积")
        if energy > allowed + TOLERANCE:
            checks["energy_reserve"] = False
            errors.append(f"{batch_id} 超过返航安全余量")
        if distance > range_m + TOLERANCE:
            checks["range_constraint"] = False
            errors.append(f"{batch_id} 超过当前载荷等效航程")

        values = [
            mass, volume, time_s, handling_s, work_time_s, energy, soc_percent, safe_payload,
            safe_payload - mass, uav.volume_capacity_m3 - volume,
            distance, range_m, range_m - distance, allowed, allowed - energy,
            item["source_mass_kg"], item["source_volume_m3"], item["source_time_s"],
            item["source_energy_kwh"], item["source_soc_percent"],
        ]
        if not all(math.isfinite(float(value)) for value in values):
            checks["finite_units"] = False
            errors.append(f"{batch_id} 存在非有限数值")

        detail_rows.append(
            {
                "batch_id": batch_id,
                "area_id": area_id,
                "uav_type": type_id,
                "box_ids_normalized": ";".join(box_ids),
                "box_count": len(box_ids),
                "mass_kg": _json_number(mass),
                "volume_m3": _json_number(volume),
                "time_s": _json_number(time_s),
                "handling_time_s": _json_number(handling_s),
                "work_time_s": _json_number(work_time_s),
                "energy_kwh": _json_number(energy),
                "soc_percent": _json_number(soc_percent),
                "source_mass_kg": item["source_mass_kg"],
                "source_volume_m3": item["source_volume_m3"],
                "source_time_s": item["source_time_s"],
                "source_energy_kwh": item["source_energy_kwh"],
                "source_soc_percent": item["source_soc_percent"],
                "mass_diff_kg": _json_number(mass - item["source_mass_kg"]),
                "volume_diff_m3": _json_number(volume - item["source_volume_m3"]),
                "time_diff_s": _json_number(time_s - item["source_time_s"]),
                "energy_diff_kwh": _json_number(energy - item["source_energy_kwh"]),
                "soc_diff_percentage_points": _json_number(soc_percent - item["source_soc_percent"]),
                "safe_payload_kg": _json_number(safe_payload),
                "mass_slack_kg": _json_number(safe_payload - mass),
                "volume_slack_m3": _json_number(uav.volume_capacity_m3 - volume),
                "round_trip_distance_m": _json_number(distance),
                "equivalent_range_m": _json_number(range_m),
                "range_slack_m": _json_number(range_m - distance),
                "allowed_energy_kwh": _json_number(allowed),
                "energy_slack_kwh": _json_number(allowed - energy),
                "source_mass_volume_match": "通过" if mass_match and volume_match else "未通过",
                "constraint_passed": "通过" if all((mass <= safe_payload + TOLERANCE,
                                                       volume <= uav.volume_capacity_m3 + TOLERANCE,
                                                       energy <= allowed + TOLERANCE,
                                                       distance <= range_m + TOLERANCE,
                                                       box_areas == {area_id})) else "未通过",
                "备注": "统一模型重算；原表值仅作对照",
            }
        )

    expected_ids = set(box_by_id)
    counts = Counter(seen)
    if set(counts) != expected_ids or any(count != 1 for count in counts.values()):
        checks["all_boxes_present_once"] = False
        missing = sorted(expected_ids - set(counts))
        duplicated = sorted(box_id for box_id, count in counts.items() if count > 1)
        errors.append(f"货箱覆盖不满足一次性：缺失={missing}，重复={duplicated}")

    checks["n_rows_equals_sorties"] = len(detail_rows) == len(rows)
    if not checks["n_rows_equals_sorties"]:
        errors.append("架次行数解析异常")
    validation = {
        "passed": not errors,
        "checks": checks,
        "errors": errors,
        "n_boxes": len(expected_ids),
        "n_batches": len(detail_rows),
    }

    source_energy = sum(row["source_energy_kwh"] for row in detail_rows)
    source_time = sum(row["source_time_s"] for row in detail_rows)
    recomputed_energy = sum(row["energy_kwh"] for row in detail_rows)
    recomputed_time = sum(row["time_s"] for row in detail_rows)
    recomputed_handling_time = sum(row["handling_time_s"] for row in detail_rows)
    recomputed_work_time = sum(row["work_time_s"] for row in detail_rows)
    metrics = {
        "source_reported": {
            "sorties": len(detail_rows),
            "energy_kwh": _json_number(source_energy),
            "time_s": _json_number(source_time),
        },
        "recomputed": {
            "sorties": len(detail_rows),
            "energy_kwh": _json_number(recomputed_energy),
            "time_s": _json_number(recomputed_time),
            "handling_time_s": _json_number(recomputed_handling_time),
            "work_time_s": _json_number(recomputed_work_time),
        },
        "delta": {
            "energy_kwh": _json_number(recomputed_energy - source_energy),
            "time_s": _json_number(recomputed_time - source_time),
        },
    }
    model = {
        "name": "reference-model-A-v1",
        "provider_module": energy_provider_module,
        "provider_source": provider.source,
        "scope": "固定服务区内 O01→S_i→O01 单点往返；外部方案给定的箱体分组和机型不重新优化",
        "time_definition": "总作业时间 = 往返飞行时间 + 工位固定准备时间 + 每箱装载时间 + 接收点基础交接时间 + 每箱增加交接时间",
        "time_source": "运输无人机数据.xlsx；Q1 题面/附录 2 的飞行时间规则",
        "caveat": "题面只给出 E=E_hor+E_up，未展开子项；reference-model-A-v1 是项目登记的参数化参考模型，重算值为条件性结果。",
    }

    headers = [
        "batch_id", "area_id", "uav_type", "box_ids_normalized", "box_count", "mass_kg", "volume_m3",
        "time_s", "handling_time_s", "work_time_s", "energy_kwh", "soc_percent", "source_mass_kg", "source_volume_m3", "source_time_s",
        "source_energy_kwh", "source_soc_percent", "mass_diff_kg", "volume_diff_m3", "time_diff_s",
        "energy_diff_kwh", "soc_diff_percentage_points", "safe_payload_kg", "mass_slack_kg", "volume_slack_m3",
        "round_trip_distance_m", "equivalent_range_m", "range_slack_m", "allowed_energy_kwh", "energy_slack_kwh",
        "source_mass_volume_match", "constraint_passed", "备注",
    ]
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "q1_对照方案_统一模型.csv", headers, detail_rows)
    _write_workbook(excel_output, solution_path, detail_rows, validation, metrics, model)

    result = {
        "status": "complete" if validation["passed"] else "failed_validation",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_solution": {
            "path": str(solution_path),
            "sha256": _sha256(solution_path),
            "sheet": SOLUTION_SHEET if SOLUTION_SHEET in sheetnames else "Q1 candidate",
            "sheetnames": sheetnames,
            "original_is_unchanged": True,
        },
        "data_dir": str(data_dir),
        "model": model,
        "metrics": metrics,
        "validation": validation,
        "outputs": {
            "csv": str((output_dir / "q1_对照方案_统一模型.csv").resolve()),
            "xlsx": str(excel_output),
        },
        "per_batch": detail_rows,
    }
    (output_dir / "q1_对照方案_统一模型.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="按统一参考模型重算外部 Q1 组批方案")
    parser.add_argument(
        "--solution",
        type=Path,
        default=Path("data/raw/D题/对照方案_问题1_结果提交.xlsx"),
        help="外部 Q1 方案的只读 XLSX 快照",
    )
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--output-dir", type=Path, default=Path("results"))
    parser.add_argument(
        "--excel-output",
        type=Path,
        default=Path("results/问题1_对照方案_统一模型.xlsx"),
    )
    parser.add_argument(
        "--energy-provider",
        default="src.problem_energy_provider",
        help="显式能耗适配器模块",
    )
    args = parser.parse_args(argv)
    result = recompute(
        args.solution,
        args.data_dir,
        args.output_dir,
        args.excel_output,
        args.energy_provider,
    )
    metrics = result["metrics"]
    print(
        "recomputed "
        f"{metrics['recomputed']['sorties']} sorties, "
        f"{metrics['recomputed']['energy_kwh']:.12f} kWh, "
        f"flight={metrics['recomputed']['time_s']:.12f} s, "
        f"handling={metrics['recomputed']['handling_time_s']:.12f} s, "
        f"work={metrics['recomputed']['work_time_s']:.12f} s"
    )
    print(f"validation={result['validation']['passed']}")
    return 0 if result["validation"]["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
